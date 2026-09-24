import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import unicodedata
import ctypes
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Optional, cast

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import fitz
import pandas as pd
from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# Rutas y configuracion
# ---------------------------------------------------------------------------

APP_NAME = "Busqueda_V1"
APP_USER_MODEL_ID = "FranciscoNieto.BusquedaV1.Desktop"
CREADO_POR = "Francisco Nieto"
BASE_DIR = Path(__file__).resolve().parent
ASSET_DIR = Path(getattr(sys, "_MEIPASS", BASE_DIR))


def obtener_ruta_configuracion() -> Path:
    carpeta_appdata = os.environ.get("APPDATA")
    if carpeta_appdata:
        carpeta_config = Path(carpeta_appdata) / APP_NAME
    else:
        carpeta_config = Path.home() / APP_NAME
    carpeta_config.mkdir(parents=True, exist_ok=True)
    return carpeta_config / "buscar_pdf_xml_excel_config.json"


CONFIG_PATH = obtener_ruta_configuracion()


def clave_cache_archivo(ruta: str) -> Optional[str]:
    try:
        path = Path(ruta).resolve()
        stat = path.stat()
    except OSError:
        return None
    return f"{path}|{stat.st_size}|{stat.st_mtime_ns}"


def clave_cache_carpeta(ruta: str) -> Optional[str]:
    try:
        path = Path(ruta).resolve()
        stat = path.stat()
    except OSError:
        return None
    return f"{path}|{stat.st_mtime_ns}"

CONFIG_DEFAULT: dict = {
    "ruta_excel": "",
    "nombre_hoja": "",
    "tipo_busqueda": "PDF y XML",
    "filtrar_mes": False,
    "columna_mes": "",
    "mes_filtro": "",
    "usar_factura": True,
    "usar_proveedores": False,
    "usar_complemento": True,
    "usar_pedimento": True,
    "usar_proforma": False,
    "usar_factura_bodega": False,
    "usar_pedimento_fb": False,
    "usar_po": True,
    "reprocesar_pedimento": False,
    "reprocesar_po": False,
    "busqueda_factura": "E",
    "resultado_factura_pdf": "T",
    "resultado_factura_xml": "U",
    "ruta_documentos_factura": "",
    "ruta_destino_factura": "",
    "busqueda_proveedores_uuid": "UUID",
    "resultado_proveedores_pdf": "Z",
    "resultado_proveedores_xml": "AA",
    "columna_proveedores_nombre_pdf": "Ordenes de compra",
    "ruta_documentos_proveedores": "",
    "ruta_destino_proveedores": "",
    "busqueda_complemento": "F",
    "resultado_complemento_pdf": "V",
    "resultado_complemento_xml": "W",
    "ruta_documentos_complemento": "",
    "ruta_destino_complemento": "",
    "busqueda_pedimento": "G",
    "resultado_pedimento_pdf": "X",
    "uuid_pedimento": "",
    "resultado_pedimento_uuid": "",
    "ruta_documentos_pedimento": "",
    "ruta_destino_pedimento": "",
    "usar_b_pedimentos": False,
    "bp_uuid": "UUID",
    "bp_folio": "Folio",
    "bp_cliente": "",
    "bp_resultado_pedimento": "Pedimento encontrado",
    "bp_resultado_proforma": "Proforma encontrada",
    "bp_ruta_pedimentos": "",
    "bp_ruta_documentos": "",
    "bp_filtrar_carpetas": "",
    "busqueda_proforma": "",
    "resultado_proforma_pdf": "",
    "ruta_documentos_proforma": "",
    "ruta_destino_proforma": "",
    "busqueda_factura_bodega": "",
    "resultado_factura_bodega_pdf": "",
    "ruta_documentos_factura_bodega": "",
    "ruta_destino_factura_bodega": "",
    "busqueda_pedimento_fb": "",
    "resultado_pedimento_fb_pdf": "",
    "ruta_documentos_pedimento_fb": "",
    "ruta_destino_pedimento_fb": "",
    "busqueda_po": "H",
    "resultado_po_pdf": "Y",
    "nombre_cliente_po": "",
    "ruta_documentos_po": "",
    "ruta_destino_po": "",
    "ocr_po_habilitado": True,
}


# ---------------------------------------------------------------------------
# Dataclass para cada bloque
# ---------------------------------------------------------------------------

@dataclass
class BloqueConfig:
    nombre: str
    activo: bool
    col_busqueda: str
    col_resultado_pdf: str
    ruta_documentos: str
    ruta_destino: str
    tiene_xml: bool = False
    col_resultado_xml: Optional[str] = None
    es_po: bool = False
    es_proveedor: bool = False
    col_nombre_pdf: str = ""
    folios_multiples: bool = False


# ---------------------------------------------------------------------------
# Persistencia
# ---------------------------------------------------------------------------

def cargar_configuracion() -> dict:
    config = dict(CONFIG_DEFAULT)
    if CONFIG_PATH.is_file():
        try:
            datos = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
            if isinstance(datos, dict):
                config.update(datos)
        except (OSError, json.JSONDecodeError):
            pass
    return config


def guardar_configuracion(config: dict) -> None:
    CONFIG_PATH.write_text(
        json.dumps(config, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def formatear_segundos(segundos: Optional[float]) -> str:
    if segundos is None or segundos < 0:
        return "--:--"
    total = int(segundos)
    minutos, seg = divmod(total, 60)
    horas, minutos = divmod(minutos, 60)
    if horas:
        return f"{horas:02d}:{minutos:02d}:{seg:02d}"
    return f"{minutos:02d}:{seg:02d}"


def normalizar_texto(valor) -> str:
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    texto = str(valor).strip()
    return texto[:-2] if texto.endswith(".0") else texto


def normalizar_para_busqueda(valor) -> str:
    texto = normalizar_texto(valor).casefold()
    texto = "".join(
        caracter for caracter in unicodedata.normalize("NFKD", texto)
        if not unicodedata.combining(caracter)
    )
    return re.sub(r"[^a-z0-9]", "", texto)


def clave_folio_busqueda(folio: str) -> str:
    return normalizar_para_busqueda(folio) or normalizar_texto(folio).casefold()


MESES_ES = {
    1: ("enero", "ene"),
    2: ("febrero", "feb"),
    3: ("marzo", "mar"),
    4: ("abril", "abr"),
    5: ("mayo", "may"),
    6: ("junio", "jun"),
    7: ("julio", "jul"),
    8: ("agosto", "ago"),
    9: ("septiembre", "sep"),
    10: ("octubre", "oct"),
    11: ("noviembre", "nov"),
    12: ("diciembre", "dic"),
}

MESES_POR_TEXTO = {
    nombre: numero
    for numero, nombres in MESES_ES.items()
    for nombre in nombres
}

UUID_RE = re.compile(
    r"\b[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-"
    r"[0-9a-fA-F]{4}-[0-9a-fA-F]{12}\b"
)

PEDIMENTO_FORMATTED_RE = re.compile(
    r"(?<!\d)(\d{2})[\s.\-]*(\d{2})[\s.\-]*(\d{4})[\s.\-]*(\d{7})(?!\d)"
)

PEDIMENTO_LABEL_RE = re.compile(
    r"(?:pedimento|ped\.?)\D{0,40}(\d{2}[\s.\-]*\d{2}[\s.\-]*\d{4}[\s.\-]*\d{7})",
    re.IGNORECASE,
)

PEDIMENTO_CGR_RE = re.compile(
    r"(?<!\d)(\d{4})[\s._\-]+(\d{3,4})[\s._\-]+(\d{6,7})(?!\d)"
)

PEDIMENTO_CGR_SPACED_RE = re.compile(
    r"(?<!\d)(\d{2})[\s._\-]+(\d{2})[\s._\-]+(\d{3,4})[\s._\-]+(\d{6,7})(?!\d)"
)

PROFORMA_RE = re.compile(r"\bCGR\s*-\s*\d{4}\s*-\s*\d+\b", re.IGNORECASE)


def normalizar_uuid(valor) -> str:
    texto = normalizar_texto(valor)
    match = UUID_RE.search(texto)
    if match:
        return match.group(0).upper()
    solo_hex = re.sub(r"[^0-9a-fA-F]", "", texto)
    if len(solo_hex) == 32:
        return (
            f"{solo_hex[0:8]}-{solo_hex[8:12]}-{solo_hex[12:16]}-"
            f"{solo_hex[16:20]}-{solo_hex[20:32]}"
        ).upper()
    return texto.upper()


def normalizar_pedimento_extraido(valor: str) -> str:
    digitos = re.sub(r"\D", "", valor)
    if len(digitos) == 15:
        return f"{digitos[0:2]} {digitos[2:4]} {digitos[4:8]} {digitos[8:15]}"
    return normalizar_texto(valor)


def extraer_pedimento_de_texto(texto: str) -> str:
    match = PEDIMENTO_LABEL_RE.search(texto)
    if match:
        return normalizar_pedimento_extraido(match.group(1))

    match = PEDIMENTO_CGR_RE.search(texto)
    if match:
        return " ".join(match.groups())

    match = PEDIMENTO_CGR_SPACED_RE.search(texto)
    if match:
        return f"{match.group(1)}{match.group(2)} {match.group(3)} {match.group(4)}"

    for match in PEDIMENTO_FORMATTED_RE.finditer(texto):
        candidato = "".join(match.groups())
        if len(candidato) == 15:
            return normalizar_pedimento_extraido(candidato)
    return ""


def normalizar_proforma(valor: str) -> str:
    return re.sub(r"\s*-\s*", "-", normalizar_texto(valor).upper())


def extraer_proforma_de_texto(texto: str) -> str:
    match = PROFORMA_RE.search(texto)
    return normalizar_proforma(match.group(0)) if match else ""


def parsear_filtros_ruta(filtros: str) -> list[str]:
    return [filtro.strip().casefold() for filtro in normalizar_texto(filtros).split(",") if filtro.strip()]


def filtrar_archivos_por_ruta(archivos: list, filtros: str) -> list:
    filtros_ruta = parsear_filtros_ruta(filtros)
    if not filtros_ruta:
        return archivos
    return [
        archivo for archivo in archivos
        if any(filtro in archivo[2].casefold() for filtro in filtros_ruta)
    ]


def claves_mes(valor) -> set[str]:
    try:
        if pd.isna(valor):
            return set()
    except (TypeError, ValueError):
        pass

    numero_mes = None
    if isinstance(valor, (datetime, date)):
        numero_mes = valor.month
    else:
        texto = normalizar_texto(valor)
        texto_simple = "".join(
            caracter for caracter in unicodedata.normalize("NFKD", texto.casefold())
            if not unicodedata.combining(caracter)
        ).strip()

        if re.fullmatch(r"\d{1,2}", texto_simple):
            posible_mes = int(texto_simple)
            if 1 <= posible_mes <= 12:
                numero_mes = posible_mes

        if numero_mes is None:
            for nombre_mes, numero in MESES_POR_TEXTO.items():
                if re.search(rf"\b{re.escape(nombre_mes)}\b", texto_simple):
                    numero_mes = numero
                    break

        if numero_mes is None:
            fecha = pd.to_datetime(texto, errors="coerce", dayfirst=True)
            if not pd.isna(fecha):
                numero_mes = int(fecha.month)

    if numero_mes is None:
        return set()

    nombre, abreviado = MESES_ES[numero_mes]
    return {str(numero_mes), f"{numero_mes:02d}", nombre, abreviado}


def coincide_mes(valor_excel, mes_objetivo: str) -> bool:
    claves_objetivo = claves_mes(mes_objetivo)
    if not claves_objetivo:
        return False
    return bool(claves_mes(valor_excel) & claves_objetivo)


def separar_folios(valor: str) -> list[str]:
    folios = []
    vistos = set()
    for parte in normalizar_texto(valor).split(","):
        folio = normalizar_texto(parte)
        clave = clave_folio_busqueda(folio)
        if folio and clave and clave not in vistos:
            folios.append(folio)
            vistos.add(clave)
    return folios


def variantes_busqueda_pedimento(folio: str) -> list[str]:
    folio = normalizar_texto(folio)
    variantes = []
    vistos = set()

    def agregar(valor: str) -> None:
        valor = normalizar_texto(valor)
        clave = clave_folio_busqueda(valor)
        if valor and clave and clave not in vistos:
            variantes.append(valor)
            vistos.add(clave)

    agregar(folio)
    partes = [normalizar_texto(parte) for parte in re.split(r"[-_\s]+", folio) if normalizar_texto(parte)]
    if len(partes) > 1:
        ultimo = partes[-1]
        if len(normalizar_para_busqueda(ultimo)) >= 5:
            agregar(ultimo)
    return variantes


def partes_pedimento(folio: str) -> list[str]:
    return [
        normalizar_texto(parte)
        for parte in re.split(r"[-_\s]+", normalizar_texto(folio))
        if normalizar_texto(parte)
    ]


def coincide_pedimento_en_texto(folio: str, texto: str) -> bool:
    texto_normalizado = normalizar_para_busqueda(texto)
    if not texto_normalizado:
        return False

    for variante in variantes_busqueda_pedimento(folio):
        variante_normalizada = normalizar_para_busqueda(variante)
        if variante_normalizada and variante_normalizada in texto_normalizado:
            return True

    partes = partes_pedimento(folio)
    if len(partes) >= 2:
        primero = normalizar_para_busqueda(partes[0])
        ultimo = normalizar_para_busqueda(partes[-1])
        if primero and len(ultimo) >= 5:
            inicio = texto_normalizado.find(primero)
            fin = texto_normalizado.find(ultimo, inicio + len(primero)) if inicio >= 0 else -1
            if fin >= 0:
                return True

    return False


def resultado_pdf_no(valor) -> bool:
    return normalizar_texto(valor).casefold().startswith("pdf no")


def resultado_folios_multiples(encontrados: list[str], faltantes: list[str]) -> str:
    if encontrados and not faltantes:
        return "PDF OK"
    if encontrados and faltantes:
        return f"PDF OK, pero falta: {', '.join(faltantes)}"
    return f"PDF NO: {', '.join(faltantes)}"


def es_pdf_completo(nombre: str) -> bool:
    nombre_normalizado = normalizar_para_busqueda(nombre)
    return "completo" in nombre_normalizado


def letra_a_indice_columna(letra_columna: str) -> int:
    letra = letra_columna.strip().upper()
    if not letra.isalpha():
        raise ValueError(f"La columna '{letra_columna}' no es valida.")
    indice = 0
    for caracter in letra:
        indice = indice * 26 + (ord(caracter) - ord("A") + 1)
    return indice - 1


def indice_columna_por_letra_o_encabezado(df: pd.DataFrame, identificador: str) -> Optional[int]:
    identificador = normalizar_texto(identificador)
    if not identificador:
        return None

    if identificador.isalpha() and len(identificador) <= 3:
        idx = letra_a_indice_columna(identificador)
        if idx < len(df.columns):
            return idx

    clave = normalizar_para_busqueda(identificador)
    for idx, columna in enumerate(df.columns):
        if normalizar_para_busqueda(columna) == clave:
            return idx
    return None


def indice_hoja_por_letra_o_encabezado(hoja, identificador: str, crear: bool = False) -> int:
    identificador = normalizar_texto(identificador)
    if identificador.isalpha() and len(identificador) <= 3:
        return letra_a_indice_columna(identificador) + 1

    clave = normalizar_para_busqueda(identificador)
    for celda in hoja[1]:
        if normalizar_para_busqueda(celda.value) == clave:
            return celda.column

    if crear:
        columna = hoja.max_column + 1
        hoja.cell(row=1, column=columna).value = identificador
        return columna

    raise ValueError(f"No encontre la columna '{identificador}' en los encabezados.")


# ---------------------------------------------------------------------------
# Indices y busqueda de archivos
# ---------------------------------------------------------------------------

def construir_indices(
    carpeta_base: str,
    cache_indices: Optional[dict[str, tuple[list, list]]] = None,
) -> tuple[list, list]:
    clave_cache = clave_cache_carpeta(carpeta_base)
    if cache_indices is not None and clave_cache and clave_cache in cache_indices:
        pdfs_cache, xmls_cache = cache_indices[clave_cache]
        return list(pdfs_cache), list(xmls_cache)

    pdfs, xmls = [], []
    for raiz, _, archivos in os.walk(carpeta_base):
        for archivo in archivos:
            nombre = archivo.casefold()
            nombre_normalizado = normalizar_para_busqueda(Path(archivo).stem)
            ruta_completa = str(Path(raiz) / archivo)
            if nombre.endswith(".pdf"):
                pdfs.append((nombre, nombre_normalizado, ruta_completa))
            elif nombre.endswith(".xml"):
                xmls.append((nombre, nombre_normalizado, ruta_completa))

    if cache_indices is not None and clave_cache:
        cache_indices[clave_cache] = (list(pdfs), list(xmls))
    return pdfs, xmls


def existe_documento(folio: str, archivos: list) -> Optional[str]:
    folio_lower = normalizar_texto(folio).casefold()
    folio_normalizado = normalizar_para_busqueda(folio)
    if not folio_lower:
        return None

    for nombre, nombre_normalizado, ruta in archivos:
        if folio_lower in nombre or (
            folio_normalizado and folio_normalizado in nombre_normalizado
        ):
            return ruta
    return None


def existe_documento_pedimento(folio: str, archivos: list) -> Optional[str]:
    if not normalizar_texto(folio):
        return None

    for nombre, nombre_normalizado, ruta in archivos:
        if coincide_pedimento_en_texto(folio, nombre_normalizado or nombre):
            return ruta
    return None


def copiar_documento(ruta_origen: str, carpeta_destino: Path) -> None:
    origen = Path(ruta_origen)
    destino = carpeta_destino / origen.name
    try:
        if origen.resolve() == destino.resolve():
            return
    except OSError:
        pass
    shutil.copy2(origen, destino)


def nombre_archivo_seguro(nombre: str) -> str:
    nombre = normalizar_texto(nombre)
    nombre = re.sub(r'[<>:"/\\|?*]+', "_", nombre)
    nombre = re.sub(r"\s+", " ", nombre).strip(" .")
    return nombre or "documento"


def ruta_destino_disponible(carpeta_destino: Path, nombre_base: str, extension: str) -> Path:
    destino = carpeta_destino / f"{nombre_base}{extension}"
    contador = 2
    while destino.exists():
        destino = carpeta_destino / f"{nombre_base}_{contador}{extension}"
        contador += 1
    return destino


def copiar_documento_como(ruta_origen: str, carpeta_destino: Path, nombre_base: str) -> Path:
    origen = Path(ruta_origen)
    nombre_seguro = nombre_archivo_seguro(nombre_base)
    extension = origen.suffix or ".pdf"
    destino = ruta_destino_disponible(carpeta_destino, nombre_seguro, extension)
    try:
        if origen.resolve() == destino.resolve():
            return destino
    except OSError:
        pass
    shutil.copy2(origen, destino)
    return destino


# ---------------------------------------------------------------------------
# OCR y busqueda especial para PO
# ---------------------------------------------------------------------------

def ocr_imagen_windows(ruta_imagen: Path) -> str:
    script = r"""
Add-Type -AssemblyName System.Runtime.WindowsRuntime
[Windows.Storage.StorageFile, Windows.Storage, ContentType = WindowsRuntime] > $null
[Windows.Storage.FileAccessMode, Windows.Storage, ContentType = WindowsRuntime] > $null
[Windows.Graphics.Imaging.BitmapDecoder, Windows.Graphics.Imaging, ContentType = WindowsRuntime] > $null
[Windows.Media.Ocr.OcrEngine, Windows.Media.Ocr, ContentType = WindowsRuntime] > $null

function Await($operation) {
    [System.WindowsRuntimeSystemExtensions]::AsTask($operation).GetAwaiter().GetResult()
}

[Console]::OutputEncoding = [System.Text.Encoding]::UTF8
$path = $env:OCR_IMAGE_PATH
$file = Await([Windows.Storage.StorageFile]::GetFileFromPathAsync($path))
$stream = Await($file.OpenAsync([Windows.Storage.FileAccessMode]::Read))
$decoder = Await([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($stream))
$bitmap = Await($decoder.GetSoftwareBitmapAsync())
$engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages()
$result = Await($engine.RecognizeAsync($bitmap))
Write-Output $result.Text
    """
    entorno = os.environ.copy()
    entorno["OCR_IMAGE_PATH"] = str(ruta_imagen)
    startupinfo = None
    creationflags = 0

    if os.name == "nt":
        startupinfo = subprocess.STARTUPINFO()
        startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0)

    try:
        resultado = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
            env=entorno,
            timeout=90,
            startupinfo=startupinfo,
            creationflags=creationflags,
        )
        if resultado.returncode != 0:
            return ""
        return resultado.stdout.strip()
    except subprocess.TimeoutExpired:
        return ""
    except Exception:
        return ""


def extraer_texto_pdf_po_sin_ocr(
    ruta_pdf: str,
    cache_texto: dict[str, str],
) -> str:
    clave_cache = clave_cache_archivo(ruta_pdf) or ruta_pdf
    if clave_cache in cache_texto:
        return cache_texto[clave_cache]

    partes: list[str] = []
    try:
        with fitz.open(ruta_pdf) as documento:
            for pagina in documento:
                texto = cast(str, pagina.get_text("text") or "").casefold()
                if texto:
                    partes.append(texto)
    except Exception:
        cache_texto[clave_cache] = ""
        return ""

    contenido = "\n".join(partes)
    cache_texto[clave_cache] = contenido
    return contenido


def extraer_ocr_pdf_po(ruta_pdf: str, cache_ocr: dict[str, str]) -> str:
    clave_cache = clave_cache_archivo(ruta_pdf) or ruta_pdf
    if clave_cache in cache_ocr:
        return cache_ocr[clave_cache]

    partes: list[str] = []
    try:
        with fitz.open(ruta_pdf) as documento:
            for pagina in documento:
                pix = pagina.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                temporal = None
                try:
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as f:
                        temporal = Path(f.name)
                    pix.save(temporal)
                    texto_ocr = ocr_imagen_windows(temporal).casefold()
                    if texto_ocr:
                        partes.append(texto_ocr)
                finally:
                    if temporal and temporal.exists():
                        temporal.unlink(missing_ok=True)
    except Exception:
        cache_ocr[clave_cache] = ""
        return ""

    contenido = "\n".join(partes)
    cache_ocr[clave_cache] = contenido
    return contenido


def preparar_busqueda_proveedores(
    uuids: list[str],
    archivos_pdf: list,
    cache_texto: dict[str, str],
    cache_ocr: dict[str, str],
    log=None,
) -> dict[str, dict]:
    pendientes = {
        normalizar_para_busqueda(uuid): normalizar_texto(uuid)
        for uuid in uuids
        if normalizar_para_busqueda(uuid)
    }
    resultados = {
        clave: {"ruta": None, "origen": "", "uuid": uuid}
        for clave, uuid in pendientes.items()
    }
    total_pdfs = len(archivos_pdf)
    if not pendientes:
        return resultados

    def registrar(clave: str, ruta_pdf: str, origen: str) -> None:
        if clave in pendientes:
            resultados[clave] = {
                "ruta": ruta_pdf,
                "origen": origen,
                "uuid": pendientes[clave],
            }
            pendientes.pop(clave, None)

    if log:
        log(
            f"Proveedores etapa 1/3: buscando {len(pendientes)} UUIDs en nombres de "
            f"{total_pdfs} PDFs..."
        )
    for numero_pdf, (_, nombre_normalizado, ruta_pdf) in enumerate(archivos_pdf, start=1):
        for clave in list(pendientes):
            if clave in nombre_normalizado:
                registrar(clave, ruta_pdf, "nombre del archivo")
        if log and (numero_pdf == 1 or numero_pdf % 50 == 0 or numero_pdf == total_pdfs):
            log(
                f"Proveedores etapa 1/3: revisados {numero_pdf}/{total_pdfs} PDFs; "
                f"faltan {len(pendientes)} UUIDs."
            )
        if not pendientes:
            return resultados

    if log:
        log(
            f"Proveedores etapa 2/3: buscando {len(pendientes)} UUIDs en texto interno..."
        )
    for numero_pdf, (_, _, ruta_pdf) in enumerate(archivos_pdf, start=1):
        contenido = normalizar_para_busqueda(
            extraer_texto_pdf_po_sin_ocr(ruta_pdf, cache_texto)
        )
        if contenido:
            for clave in list(pendientes):
                if clave in contenido:
                    registrar(clave, ruta_pdf, "texto del PDF")
        if log and (numero_pdf == 1 or numero_pdf % 25 == 0 or numero_pdf == total_pdfs):
            log(
                f"Proveedores etapa 2/3: revisados {numero_pdf}/{total_pdfs} PDFs; "
                f"faltan {len(pendientes)} UUIDs."
            )
        if not pendientes:
            return resultados

    if log:
        log(
            f"Proveedores etapa 3/3: buscando {len(pendientes)} UUIDs por OCR. "
            "Esta etapa puede tardar varios minutos si hay muchos PDFs escaneados."
        )
    for numero_pdf, (_, _, ruta_pdf) in enumerate(archivos_pdf, start=1):
        contenido = normalizar_para_busqueda(extraer_ocr_pdf_po(ruta_pdf, cache_ocr))
        if contenido:
            for clave in list(pendientes):
                if clave in contenido:
                    registrar(clave, ruta_pdf, "OCR texto/imagen")
        if log and (numero_pdf == 1 or numero_pdf % 10 == 0 or numero_pdf == total_pdfs):
            log(
                f"Proveedores etapa 3/3: OCR {numero_pdf}/{total_pdfs} PDFs; "
                f"faltan {len(pendientes)} UUIDs."
            )
        if not pendientes:
            return resultados

    if log and pendientes:
        log(f"Proveedores: {len(pendientes)} UUIDs no fueron encontrados.")
    return resultados


def es_pdf_candidato_po(nombre: str) -> bool:
    nombre_normalizado = normalizar_para_busqueda(nombre)
    palabras_po = (
        "po",
        "oc",
        "ordencompra",
        "ordendecompra",
        "purchaseorder",
        "schedulingagreement",
    )
    return any(palabra in nombre_normalizado for palabra in palabras_po)


def puntaje_cliente_en_ruta(cliente: str, ruta_pdf: str) -> int:
    cliente_normalizado = normalizar_para_busqueda(cliente)
    if not cliente_normalizado:
        return 0

    ruta = Path(ruta_pdf)
    partes_ruta = [normalizar_para_busqueda(parte) for parte in ruta.parts[:-1]]
    nombre_archivo = normalizar_para_busqueda(ruta.stem)

    if any(cliente_normalizado == parte for parte in partes_ruta):
        return 4
    if any(cliente_normalizado in parte or parte in cliente_normalizado for parte in partes_ruta if parte):
        return 3
    if cliente_normalizado in nombre_archivo:
        return 2
    tokens_cliente = [token for token in re.split(r"[^a-z0-9]+", normalizar_texto(cliente).casefold()) if len(token) >= 4]
    if tokens_cliente and any(token in nombre_archivo for token in tokens_cliente):
        return 1
    return 0


def ordenar_pdfs_para_po(archivos_pdf: list, clientes_por_clave: Optional[dict[str, str]] = None) -> list:
    clientes = list((clientes_por_clave or {}).values())

    def prioridad(archivo):
        nombre, _, ruta = archivo
        mejor_cliente = max((puntaje_cliente_en_ruta(cliente, ruta) for cliente in clientes), default=0)
        es_po = 1 if es_pdf_candidato_po(nombre) else 0
        return (-mejor_cliente, -es_po, nombre)

    return sorted(archivos_pdf, key=prioridad)


def ordenar_pdfs_para_folio_cliente(archivos_pdf: list, cliente: str) -> list:
    def prioridad(archivo):
        nombre, _, ruta = archivo
        puntaje_cliente = puntaje_cliente_en_ruta(cliente, ruta)
        es_po = 1 if es_pdf_candidato_po(nombre) else 0
        return (-puntaje_cliente, -es_po, nombre)

    return sorted(archivos_pdf, key=prioridad)


def contar_pdfs_cliente(archivos_pdf: list, cliente: str) -> int:
    if not cliente:
        return 0
    return sum(1 for _, _, ruta in archivos_pdf if puntaje_cliente_en_ruta(cliente, ruta) > 0)


def texto_pdf_requiere_ocr(contenido_texto: str) -> bool:
    texto_normalizado = normalizar_para_busqueda(contenido_texto)
    return len(texto_normalizado) < 40


def preparar_busqueda_po(
    folios: list[str],
    archivos_pdf: list,
    cache_po: dict,
    ocr_habilitado: bool,
    clientes_por_clave: Optional[dict[str, str]] = None,
    log=None,
    on_found=None,
    on_not_found=None,
) -> dict[str, dict]:
    folios_unicos: list[str] = []
    folios_por_clave: dict[str, str] = {}
    resultados: dict[str, dict] = {}
    cache_texto: dict = cache_po.setdefault("texto", {})
    cache_ocr: dict = cache_po.setdefault("ocr", {})

    for folio in folios:
        if not folio:
            continue
        clave = clave_folio_busqueda(folio)
        if not clave or clave in folios_por_clave:
            continue
        folios_por_clave[clave] = folio
        folios_unicos.append(folio)
        resultados[clave] = {"ruta": None, "origen": "", "folio": folio}

    total_folios = len(folios_unicos)
    clientes_por_clave = clientes_por_clave or {}
    archivos_pdf_po = ordenar_pdfs_para_po(archivos_pdf, clientes_por_clave)
    total_pdfs = len(archivos_pdf_po)
    total_candidatos = sum(1 for archivo in archivos_pdf_po if es_pdf_candidato_po(archivo[0]))
    total_clientes = len({cliente for cliente in clientes_por_clave.values() if cliente})
    if log:
        log(
            f"PO: {total_folios} folios unicos en memoria para buscar en {total_pdfs} PDFs "
            f"({total_candidatos} candidatos PO, {total_clientes} clientes para priorizar carpetas), incluyendo subcarpetas."
        )
        log("PO etapa 1/3: buscando primero por similitud de cliente/carpeta y folio por nombre...")

    for numero, folio in enumerate(folios_unicos, start=1):
        clave = clave_folio_busqueda(folio)
        ruta = existe_documento(folio, archivos_pdf_po)
        if ruta:
            resultados[clave] = {"ruta": ruta, "origen": "Nombre", "folio": folio}
            if log:
                log(f"PO etapa 1/3 ({numero}/{total_folios}): Folio encontrado por Nombre: '{folio}' -> {Path(ruta).name}")
            if on_found:
                on_found(clave, folio, ruta, "Nombre")
        elif log:
            log(f"PO etapa 1/3 ({numero}/{total_folios}): '{folio}' no esta en nombre, queda pendiente.")

    pendientes = {clave for clave, dato in resultados.items() if not dato["ruta"]}
    if log:
        log(f"PO etapa 2/3: buscando {len(pendientes)} folios pendientes en texto interno del PDF...")
        for clave in sorted(pendientes):
            log(f"PO etapa 2/3: buscando folio por Texto PDF: '{folios_por_clave[clave]}'")

    for clave in list(pendientes):
        folio = folios_por_clave[clave]
        cliente = clientes_por_clave.get(clave, "")
        archivos_folio = ordenar_pdfs_para_folio_cliente(archivos_pdf_po, cliente)
        total_cliente = contar_pdfs_cliente(archivos_pdf_po, cliente)

        if log and cliente:
            log(
                f"PO etapa 2/3: '{folio}' revisa primero carpetas similares a cliente "
                f"'{cliente}' ({total_cliente} PDFs priorizados)."
            )

        folio_lower = normalizar_texto(folio).casefold()
        for numero_pdf, (_, _, ruta_pdf) in enumerate(archivos_folio, start=1):
            if log and (numero_pdf == 1 or numero_pdf % 100 == 0 or numero_pdf == total_pdfs):
                log(f"PO etapa 2/3: folio '{folio}' avance Texto PDF {numero_pdf}/{total_pdfs}.")

            contenido = extraer_texto_pdf_po_sin_ocr(ruta_pdf, cache_texto)
            contenido_normalizado = normalizar_para_busqueda(contenido)

            if folio_lower in contenido or (clave and clave in contenido_normalizado):
                nombre_pdf = Path(ruta_pdf).name
                resultados[clave] = {"ruta": ruta_pdf, "origen": "Texto PDF", "folio": folio}
                pendientes.remove(clave)
                if log:
                    log(f"PO etapa 2/3: Folio encontrado por texto PDF: '{folio}' -> {nombre_pdf}")
                if on_found:
                    on_found(clave, folio, ruta_pdf, "Texto PDF")
                break

    if pendientes and ocr_habilitado:
        if log:
            log(f"PO etapa 3/3: buscando {len(pendientes)} folios pendientes por OCR texto/imagen...")
            for clave in sorted(pendientes):
                log(f"PO etapa 3/3: buscando folio por OCR texto/imagen: '{folios_por_clave[clave]}'")

        pdfs_para_ocr = [
            archivo for archivo in archivos_pdf_po
            if texto_pdf_requiere_ocr(cache_texto.get(archivo[2], ""))
        ]
        total_ocr = len(pdfs_para_ocr)
        if log:
            omitidos = total_pdfs - total_ocr
            log(
                f"PO etapa 3/3: OCR limitado a {total_ocr} PDFs con poco texto "
                f"({omitidos} PDFs omitidos por tener texto interno)."
            )

        for clave in list(pendientes):
            folio = folios_por_clave[clave]
            cliente = clientes_por_clave.get(clave, "")
            archivos_ocr_folio = ordenar_pdfs_para_folio_cliente(pdfs_para_ocr, cliente)
            total_ocr_folio = len(archivos_ocr_folio)

            if log and cliente:
                total_cliente_ocr = contar_pdfs_cliente(pdfs_para_ocr, cliente)
                log(
                    f"PO etapa 3/3: '{folio}' revisa primero OCR en carpetas similares a cliente "
                    f"'{cliente}' ({total_cliente_ocr} PDFs priorizados)."
                )

            folio_lower = normalizar_texto(folio).casefold()
            for numero_pdf, (_, _, ruta_pdf) in enumerate(archivos_ocr_folio, start=1):
                if log and (numero_pdf == 1 or numero_pdf % 25 == 0 or numero_pdf == total_ocr_folio):
                    log(f"PO etapa 3/3: folio '{folio}' avance OCR texto/imagen {numero_pdf}/{total_ocr_folio}.")

                contenido = extraer_ocr_pdf_po(ruta_pdf, cache_ocr)
                contenido_normalizado = normalizar_para_busqueda(contenido)

                if folio_lower in contenido or (clave and clave in contenido_normalizado):
                    nombre_pdf = Path(ruta_pdf).name
                    resultados[clave] = {"ruta": ruta_pdf, "origen": "OCR texto/imagen", "folio": folio}
                    pendientes.remove(clave)
                    if log:
                        log(f"PO etapa 3/3: Folio encontrado por OCR texto/imagen: '{folio}' -> {nombre_pdf}")
                    if on_found:
                        on_found(clave, folio, ruta_pdf, "OCR texto/imagen")
                    break

    if log:
        for clave in pendientes:
            log(f"PO: '{folios_por_clave[clave]}' no encontrado despues de revisar nombre, texto y OCR.")
            if on_not_found:
                on_not_found(clave, folios_por_clave[clave])
        encontrados = sum(1 for dato in resultados.values() if dato["ruta"])
        log(f"PO: busqueda terminada. Encontrados {encontrados} de {total_folios}.")

    return resultados


# ---------------------------------------------------------------------------
# Procesamiento de un folio
# ---------------------------------------------------------------------------

def procesar_folio(
    folio: str,
    archivos_pdf: list,
    archivos_xml: list,
    carpeta_destino: Path,
    folios_ya_revisados: set,
    buscar_xml: bool,
    buscar_pdf_func=None,
) -> tuple[str, str]:
    if not folio:
        return "PDF NO", ("XML NO" if buscar_xml else "")

    if folio in folios_ya_revisados:
        return "PDF OK", ("XML OK" if buscar_xml else "")

    if buscar_pdf_func:
        ruta_pdf = buscar_pdf_func(folio, archivos_pdf)
    else:
        ruta_pdf = existe_documento(folio, archivos_pdf)

    ruta_xml = existe_documento(folio, archivos_xml) if buscar_xml else None

    if ruta_pdf:
        copiar_documento(ruta_pdf, carpeta_destino)
        resultado_pdf = "PDF OK"
    else:
        resultado_pdf = "PDF NO"

    if buscar_xml and ruta_xml:
        copiar_documento(ruta_xml, carpeta_destino)
        resultado_xml = "XML OK"
    elif buscar_xml:
        resultado_xml = "XML NO"
    else:
        resultado_xml = ""

    folios_ya_revisados.add(folio)
    return resultado_pdf, resultado_xml


def buscar_folio_en_texto_pdfs_completos(
    folio: str,
    archivos_pdf: list,
    cache_texto: dict,
) -> Optional[str]:
    if not normalizar_texto(folio):
        return None

    for nombre, _, ruta_pdf in archivos_pdf:
        if not es_pdf_completo(nombre):
            continue
        contenido = extraer_texto_pdf_po_sin_ocr(ruta_pdf, cache_texto)
        if coincide_pedimento_en_texto(folio, contenido):
            return ruta_pdf
    return None


def buscar_pedimento_por_uuid_en_pdfs(
    uuid: str,
    archivos_pdf: list,
    cache_texto: dict,
) -> tuple[Optional[str], str]:
    uuid_normalizado = normalizar_uuid(uuid)
    uuid_lookup = normalizar_para_busqueda(uuid_normalizado)
    if not uuid_lookup:
        return None, ""

    for _, _, ruta_pdf in archivos_pdf:
        contenido = extraer_texto_pdf_po_sin_ocr(ruta_pdf, cache_texto)
        if uuid_lookup and uuid_lookup in normalizar_para_busqueda(contenido):
            return ruta_pdf, extraer_pedimento_de_texto(contenido)
    return None, ""


def buscar_pedimento_por_referencia_en_pdfs(
    referencia: str,
    archivos_pdf: list,
    cache_texto: dict,
) -> tuple[Optional[str], str]:
    referencia_lookup = normalizar_para_busqueda(referencia)
    if not referencia_lookup:
        return None, ""

    for _, _, ruta_pdf in archivos_pdf:
        contenido = extraer_texto_pdf_po_sin_ocr(ruta_pdf, cache_texto)
        if referencia_lookup in normalizar_para_busqueda(contenido):
            pedimento = extraer_pedimento_de_texto(contenido)
            if pedimento:
                return ruta_pdf, pedimento
    return None, ""


def buscar_proforma_por_folio_en_pdfs(
    folio: str,
    archivos_pdf: list,
    cache_texto: dict,
) -> str:
    folio_lookup = normalizar_para_busqueda(folio)
    if not folio_lookup:
        return ""

    for _, _, ruta_pdf in archivos_pdf:
        contenido = extraer_texto_pdf_po_sin_ocr(ruta_pdf, cache_texto)
        if folio_lookup in normalizar_para_busqueda(contenido):
            proforma = extraer_proforma_de_texto(contenido)
            if proforma:
                return proforma
    return ""


def procesar_folios_multiples(
    valor_folios: str,
    archivos_pdf: list,
    carpeta_destino: Path,
    folios_ya_revisados: set,
    buscar_texto_completo: bool = False,
    cache_texto: Optional[dict] = None,
) -> tuple[str, list[str], list[str], list[str]]:
    folios = separar_folios(valor_folios)
    if not folios:
        return "PDF NO", [], [], []

    encontrados = []
    faltantes = []
    repetidos = []

    for folio in folios:
        clave = clave_folio_busqueda(folio)
        if clave in folios_ya_revisados:
            encontrados.append(folio)
            repetidos.append(folio)
            continue

        encontrado_por_texto_completo = False
        if buscar_texto_completo:
            ruta_pdf = existe_documento_pedimento(folio, archivos_pdf)
        else:
            ruta_pdf = existe_documento(folio, archivos_pdf)
        if not ruta_pdf and buscar_texto_completo:
            ruta_pdf = buscar_folio_en_texto_pdfs_completos(
                folio,
                archivos_pdf,
                cache_texto if cache_texto is not None else {},
            )
            encontrado_por_texto_completo = bool(ruta_pdf)
        if ruta_pdf:
            if encontrado_por_texto_completo:
                copiar_documento_como(ruta_pdf, carpeta_destino, folio)
            else:
                copiar_documento(ruta_pdf, carpeta_destino)
            encontrados.append(folio)
            folios_ya_revisados.add(clave)
        else:
            faltantes.append(folio)

    resultado_pdf = resultado_folios_multiples(encontrados, faltantes)
    return resultado_pdf, encontrados, faltantes, repetidos


def conciliar_folio_destino(
    folio: str,
    archivos_pdf_destino: list,
    archivos_xml_destino: list,
    buscar_xml: bool,
    usar_coincidencia_pedimento: bool = False,
) -> tuple[str, str, bool, bool]:
    if not folio:
        return "PDF NO", ("XML NO" if buscar_xml else ""), False, False

    if usar_coincidencia_pedimento:
        pdf_en_destino = bool(existe_documento_pedimento(folio, archivos_pdf_destino))
    else:
        pdf_en_destino = bool(existe_documento(folio, archivos_pdf_destino))
    xml_en_destino = bool(existe_documento(folio, archivos_xml_destino)) if buscar_xml else False
    resultado_pdf = "PDF OK" if pdf_en_destino else "PDF NO"

    if buscar_xml:
        resultado_xml = "XML OK" if xml_en_destino else "XML NO"
    else:
        resultado_xml = ""

    listo = pdf_en_destino and (xml_en_destino if buscar_xml else True)
    return resultado_pdf, resultado_xml, listo, pdf_en_destino


def conciliar_folios_multiples_destino(
    valor_folios: str,
    archivos_pdf_destino: list,
    usar_coincidencia_pedimento: bool = False,
) -> tuple[str, list[str], list[str], bool]:
    folios = separar_folios(valor_folios)
    if not folios:
        return "PDF NO", [], [], False

    encontrados = []
    faltantes = []
    for folio in folios:
        if usar_coincidencia_pedimento:
            encontrado = existe_documento_pedimento(folio, archivos_pdf_destino)
        else:
            encontrado = existe_documento(folio, archivos_pdf_destino)
        if encontrado:
            encontrados.append(folio)
        else:
            faltantes.append(folio)

    resultado_pdf = resultado_folios_multiples(encontrados, faltantes)
    return resultado_pdf, encontrados, faltantes, not faltantes


# ---------------------------------------------------------------------------
# Validaciones
# ---------------------------------------------------------------------------

def validar_bloques(
    bloques: list[BloqueConfig],
    tipo_busqueda: str,
    cache_indices: Optional[dict[str, tuple[list, list]]] = None,
) -> None:
    if not any(b.activo for b in bloques):
        raise ValueError("Debes activar al menos un tipo de busqueda.")

    columnas_resultado = []
    for b in bloques:
        if not b.activo:
            continue

        if not b.col_busqueda or not b.col_resultado_pdf:
            raise ValueError(
                f"Debes indicar las columnas de busqueda y resultado PDF para {b.nombre}."
            )
        if not b.ruta_documentos:
            raise ValueError(f"Debes seleccionar la carpeta de documentos para {b.nombre}.")
        if not Path(b.ruta_documentos).is_dir():
            raise ValueError(f"La carpeta de documentos de {b.nombre} no existe.")
        if not b.ruta_destino:
            raise ValueError(f"Debes seleccionar la carpeta destino para {b.nombre}.")

        if b.es_proveedor and not b.col_nombre_pdf:
            raise ValueError(
                "Debes indicar la columna para renombrar PDFs en Proveedores."
            )
        if not b.es_proveedor:
            letra_a_indice_columna(b.col_busqueda)
        letra_a_indice_columna(b.col_resultado_pdf)
        columnas_resultado.append(b.col_resultado_pdf)

        if b.tiene_xml and tipo_busqueda == "PDF y XML":
            if not b.col_resultado_xml:
                raise ValueError(f"Debes indicar la columna de resultado XML para {b.nombre}.")
            letra_a_indice_columna(b.col_resultado_xml)
            columnas_resultado.append(b.col_resultado_xml)

        pdfs, xmls = construir_indices(b.ruta_documentos, cache_indices)
        if not pdfs:
            raise ValueError(
                f"La carpeta de documentos de {b.nombre} no contiene PDFs."
            )
        if b.tiene_xml and tipo_busqueda == "PDF y XML" and not xmls:
            raise ValueError(
                f"La carpeta de documentos de {b.nombre} no contiene XMLs."
            )

    if len(columnas_resultado) != len(set(columnas_resultado)):
        raise ValueError(
            "Hay columnas de resultado duplicadas entre bloques. "
            "Cada bloque debe usar columnas distintas."
        )


# ---------------------------------------------------------------------------
# Interfaz grafica
# ---------------------------------------------------------------------------

class Aplicacion:
    root: tk.Tk
    ruta_excel: tk.StringVar
    nombre_hoja: tk.StringVar
    tipo_busqueda: tk.StringVar
    filtrar_mes: tk.BooleanVar
    columna_mes: tk.StringVar
    mes_filtro: tk.StringVar
    usar_factura: tk.BooleanVar
    usar_proveedores: tk.BooleanVar
    usar_complemento: tk.BooleanVar
    usar_pedimento: tk.BooleanVar
    usar_proforma: tk.BooleanVar
    usar_factura_bodega: tk.BooleanVar
    usar_pedimento_fb: tk.BooleanVar
    usar_po: tk.BooleanVar
    usar_b_pedimentos: tk.BooleanVar
    reprocesar_pedimento: tk.BooleanVar
    reprocesar_po: tk.BooleanVar
    ocr_po_habilitado: tk.BooleanVar
    busqueda_factura: tk.StringVar
    resultado_factura_pdf: tk.StringVar
    resultado_factura_xml: tk.StringVar
    ruta_documentos_factura: tk.StringVar
    ruta_destino_factura: tk.StringVar
    busqueda_proveedores_uuid: tk.StringVar
    resultado_proveedores_pdf: tk.StringVar
    resultado_proveedores_xml: tk.StringVar
    columna_proveedores_nombre_pdf: tk.StringVar
    ruta_documentos_proveedores: tk.StringVar
    ruta_destino_proveedores: tk.StringVar
    busqueda_complemento: tk.StringVar
    resultado_complemento_pdf: tk.StringVar
    resultado_complemento_xml: tk.StringVar
    ruta_documentos_complemento: tk.StringVar
    ruta_destino_complemento: tk.StringVar
    busqueda_pedimento: tk.StringVar
    resultado_pedimento_pdf: tk.StringVar
    uuid_pedimento: tk.StringVar
    resultado_pedimento_uuid: tk.StringVar
    ruta_documentos_pedimento: tk.StringVar
    ruta_destino_pedimento: tk.StringVar
    bp_uuid: tk.StringVar
    bp_folio: tk.StringVar
    bp_cliente: tk.StringVar
    bp_resultado_pedimento: tk.StringVar
    bp_resultado_proforma: tk.StringVar
    bp_ruta_pedimentos: tk.StringVar
    bp_ruta_documentos: tk.StringVar
    bp_filtrar_carpetas: tk.StringVar
    busqueda_proforma: tk.StringVar
    resultado_proforma_pdf: tk.StringVar
    ruta_documentos_proforma: tk.StringVar
    ruta_destino_proforma: tk.StringVar
    busqueda_factura_bodega: tk.StringVar
    resultado_factura_bodega_pdf: tk.StringVar
    ruta_documentos_factura_bodega: tk.StringVar
    ruta_destino_factura_bodega: tk.StringVar
    busqueda_pedimento_fb: tk.StringVar
    resultado_pedimento_fb_pdf: tk.StringVar
    ruta_documentos_pedimento_fb: tk.StringVar
    ruta_destino_pedimento_fb: tk.StringVar
    busqueda_po: tk.StringVar
    resultado_po_pdf: tk.StringVar
    nombre_cliente_po: tk.StringVar
    ruta_documentos_po: tk.StringVar
    ruta_destino_po: tk.StringVar
    estado: tk.StringVar
    contador: tk.StringVar
    progreso_texto: tk.StringVar
    log_terminal: Optional[tk.Text]
    combo_hoja: ttk.Combobox
    combo_tipo: ttk.Combobox
    entrada_excel: ttk.Entry
    entrada_columna_mes: ttk.Entry
    entrada_mes_filtro: ttk.Entry
    notebook: ttk.Notebook
    boton_guardar: ttk.Button
    boton_restaurar: ttk.Button
    boton_cancelar: ttk.Button
    boton_iniciar: ttk.Button
    barra: ttk.Progressbar
    entrada_bp_uuid: ttk.Combobox
    entrada_bp_folio: ttk.Combobox
    entrada_bp_cliente: ttk.Combobox
    entrada_bp_resultado_pedimento: ttk.Combobox
    entrada_bp_resultado_proforma: ttk.Combobox
    boton_bp_columnas: ttk.Button
    entrada_bp_pedimentos: ttk.Entry
    boton_bp_pedimentos: ttk.Button
    entrada_bp_documentos: ttk.Entry
    boton_bp_documentos: ttk.Button
    entrada_bp_filtrar: ttk.Entry
    cache_indices_origen: dict[str, tuple[list, list]]
    cache_pdf_texto: dict[str, str]
    cache_pdf_ocr: dict[str, str]
    cache_po: dict
    cache_hojas_excel: dict[str, list[str]]
    cache_encabezados_excel: dict[str, list[str]]
    cache_dataframes_excel: dict[str, pd.DataFrame]

    def __init__(self):
        self._configurar_barra_tareas_windows()
        self.root = tk.Tk(className="BusquedaV1")
        self.root.title("Busqueda V1")
        self.root.geometry("1180x780")
        self.root.minsize(1040, 680)
        self.root.resizable(True, True)

        self._cancelar = threading.Event()
        self._hilo: Optional[threading.Thread] = None

        config = cargar_configuracion()
        self._init_variables(config)

        self.procesando = False
        self.controles_bloques: dict = {}
        self.cache_indices_origen: dict[str, tuple[list, list]] = {}
        self.cache_pdf_texto: dict[str, str] = {}
        self.cache_pdf_ocr: dict[str, str] = {}
        self.cache_po: dict = {}
        self.cache_hojas_excel: dict[str, list[str]] = {}
        self.cache_encabezados_excel: dict[str, list[str]] = {}
        self.cache_dataframes_excel: dict[str, pd.DataFrame] = {}

        self._configurar_estilos()
        self._crear_interfaz()

    @staticmethod
    def _configurar_barra_tareas_windows() -> None:
        if sys.platform != "win32":
            return
        try:
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(
                APP_USER_MODEL_ID
            )
        except Exception:
            pass

    def _construir_indices_origen(self, carpeta_base: str) -> tuple[list, list]:
        return construir_indices(carpeta_base, self.cache_indices_origen)

    def _leer_hojas_excel_cache(self, ruta_excel: str) -> list[str]:
        clave_cache = clave_cache_archivo(ruta_excel)
        if not clave_cache:
            return []
        if clave_cache not in self.cache_hojas_excel:
            libro = load_workbook(ruta_excel, read_only=True, data_only=False)
            try:
                self.cache_hojas_excel[clave_cache] = list(libro.sheetnames)
            finally:
                libro.close()
        return list(self.cache_hojas_excel[clave_cache])

    def _leer_encabezados_excel_cache(self, ruta_excel: str, nombre_hoja: str) -> list[str]:
        clave_archivo = clave_cache_archivo(ruta_excel)
        if not clave_archivo:
            return []
        clave_cache = f"{clave_archivo}|{nombre_hoja}"
        if clave_cache not in self.cache_encabezados_excel:
            libro = load_workbook(ruta_excel, read_only=True, data_only=False)
            try:
                if nombre_hoja not in libro.sheetnames:
                    raise ValueError(f"La hoja '{nombre_hoja}' no existe.")
                hoja = libro[nombre_hoja]
                self.cache_encabezados_excel[clave_cache] = [
                    normalizar_texto(celda.value)
                    for celda in next(hoja.iter_rows(min_row=1, max_row=1))
                    if normalizar_texto(celda.value)
                ]
            finally:
                libro.close()
        return list(self.cache_encabezados_excel[clave_cache])

    def _validar_hoja_excel(self, ruta_excel: str, nombre_hoja: str) -> list[str]:
        hojas = self._leer_hojas_excel_cache(ruta_excel)
        if not hojas:
            raise ValueError("El archivo Excel no contiene hojas disponibles.")
        if nombre_hoja not in hojas:
            raise ValueError(
                f"La hoja seleccionada '{nombre_hoja}' no existe en el archivo Excel."
            )
        return hojas

    def _leer_dataframe_excel_cache(self, ruta_excel: str, nombre_hoja: str) -> pd.DataFrame:
        self._validar_hoja_excel(ruta_excel, nombre_hoja)
        clave_archivo = clave_cache_archivo(ruta_excel)
        clave_cache = f"{clave_archivo or ruta_excel}|{nombre_hoja}"
        if clave_cache not in self.cache_dataframes_excel:
            # Pandas recibe sheet_name para cargar solo la hoja seleccionada.
            self.cache_dataframes_excel[clave_cache] = pd.read_excel(
                ruta_excel,
                sheet_name=nombre_hoja,
            )
        return self.cache_dataframes_excel[clave_cache].copy()

    def _invalidar_cache_excel(self, ruta_excel: str) -> None:
        ruta_resuelta = str(Path(ruta_excel).resolve()) if ruta_excel else ""
        if not ruta_resuelta:
            return
        for cache in (
            self.cache_hojas_excel,
            self.cache_encabezados_excel,
            self.cache_dataframes_excel,
        ):
            for clave in list(cache):
                if clave.startswith(ruta_resuelta):
                    del cache[clave]

    # ------------------------------------------------------------------
    # Variables Tk
    # ------------------------------------------------------------------

    def _init_variables(self, config: dict) -> None:
        self.ruta_excel = tk.StringVar(value=config["ruta_excel"])
        self.nombre_hoja = tk.StringVar(value=config["nombre_hoja"])
        self.tipo_busqueda = tk.StringVar(value=config["tipo_busqueda"])
        self.filtrar_mes = tk.BooleanVar(value=config.get("filtrar_mes", False))
        self.columna_mes = tk.StringVar(value=config.get("columna_mes", ""))
        self.mes_filtro = tk.StringVar(value=config.get("mes_filtro", ""))

        self.usar_factura = tk.BooleanVar(value=config["usar_factura"])
        self.usar_proveedores = tk.BooleanVar(value=config.get("usar_proveedores", False))
        self.usar_complemento = tk.BooleanVar(value=config["usar_complemento"])
        self.usar_pedimento = tk.BooleanVar(value=config["usar_pedimento"])
        self.usar_proforma = tk.BooleanVar(value=config.get("usar_proforma", False))
        self.usar_factura_bodega = tk.BooleanVar(value=config.get("usar_factura_bodega", False))
        self.usar_pedimento_fb = tk.BooleanVar(value=config.get("usar_pedimento_fb", False))
        self.usar_po = tk.BooleanVar(value=config["usar_po"])
        self.usar_b_pedimentos = tk.BooleanVar(value=config.get("usar_b_pedimentos", False))
        self.reprocesar_pedimento = tk.BooleanVar(value=config.get("reprocesar_pedimento", False))
        self.reprocesar_po = tk.BooleanVar(value=config.get("reprocesar_po", False))
        self.ocr_po_habilitado = tk.BooleanVar(value=config.get("ocr_po_habilitado", True))

        for sufijo in [
            "busqueda_factura", "resultado_factura_pdf", "resultado_factura_xml",
            "ruta_documentos_factura", "ruta_destino_factura",
            "busqueda_proveedores_uuid", "resultado_proveedores_pdf", "resultado_proveedores_xml",
            "columna_proveedores_nombre_pdf", "ruta_documentos_proveedores", "ruta_destino_proveedores",
            "busqueda_complemento", "resultado_complemento_pdf", "resultado_complemento_xml",
            "ruta_documentos_complemento", "ruta_destino_complemento",
            "busqueda_pedimento", "resultado_pedimento_pdf",
            "uuid_pedimento", "resultado_pedimento_uuid",
            "ruta_documentos_pedimento", "ruta_destino_pedimento",
            "bp_uuid", "bp_folio", "bp_resultado_pedimento", "bp_resultado_proforma",
            "bp_cliente", "bp_ruta_pedimentos", "bp_ruta_documentos", "bp_filtrar_carpetas",
            "busqueda_proforma", "resultado_proforma_pdf",
            "ruta_documentos_proforma", "ruta_destino_proforma",
            "busqueda_factura_bodega", "resultado_factura_bodega_pdf",
            "ruta_documentos_factura_bodega", "ruta_destino_factura_bodega",
            "busqueda_pedimento_fb", "resultado_pedimento_fb_pdf",
            "ruta_documentos_pedimento_fb", "ruta_destino_pedimento_fb",
            "busqueda_po", "resultado_po_pdf", "nombre_cliente_po",
            "ruta_documentos_po", "ruta_destino_po",
        ]:
            setattr(self, sufijo, tk.StringVar(value=config.get(sufijo, "")))

        self.estado = tk.StringVar(value="Listo para iniciar.")
        self.contador = tk.StringVar(value="0 de 0 filas procesadas")
        self.progreso_texto = tk.StringVar(value="0% | Transcurrido: 00:00 | Restante: --:--")
        self.log_terminal: Optional[tk.Text] = None

    # ------------------------------------------------------------------
    # Estilos
    # ------------------------------------------------------------------

    def _configurar_estilos(self) -> None:
        style = ttk.Style()
        try:
            style.theme_use("vista")
        except tk.TclError:
            try:
                style.theme_use("clam")
            except tk.TclError:
                pass
        self.root.option_add("*Font", "{Segoe UI} 10")
        style.configure(".", font=("Segoe UI", 10))
        style.configure("Primary.TButton", padding=(14, 6))
        style.configure("Danger.TButton", padding=(14, 6))
        style.configure("Panel.TFrame", relief="flat")
        style.configure("Panel.TLabel")

    # ------------------------------------------------------------------
    # Interfaz
    # ------------------------------------------------------------------

    def _crear_interfaz(self) -> None:
        contenedor = ttk.Frame(self.root, padding=12, style="Panel.TFrame")
        contenedor.pack(fill="both", expand=True, padx=10, pady=10)
        contenedor.columnconfigure(0, weight=1)
        contenedor.rowconfigure(1, weight=3)
        contenedor.rowconfigure(2, weight=3)

        superior = ttk.Frame(contenedor, style="Panel.TFrame")
        superior.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        superior.columnconfigure(0, weight=0)
        superior.columnconfigure(1, weight=1)
        superior.columnconfigure(2, weight=0)
        superior.columnconfigure(3, weight=0)
        superior.columnconfigure(4, weight=0)

        ttk.Label(superior, text="Archivo Excel", style="Panel.TLabel").grid(row=0, column=0, sticky="w", pady=(0, 4))
        self.entrada_excel = ttk.Entry(superior, textvariable=self.ruta_excel)
        self.entrada_excel.grid(
            row=1, column=0, columnspan=4, sticky="ew", padx=(0, 8)
        )
        self.entrada_excel.bind("<FocusOut>", self._al_cambiar_ruta_excel)
        self.entrada_excel.bind("<Return>", self._al_cambiar_ruta_excel)
        ttk.Button(superior, text="Examinar", command=self._seleccionar_excel).grid(row=1, column=4, sticky="e")

        ttk.Label(superior, text="Hoja de Excel", style="Panel.TLabel").grid(row=2, column=0, sticky="w", pady=(10, 4))
        self.combo_hoja = ttk.Combobox(
            superior,
            textvariable=self.nombre_hoja,
            values=[],
            state="readonly",
            width=28,
            postcommand=self._refrescar_hojas_excel,
        )
        self.combo_hoja.grid(row=3, column=0, sticky="w")
        self.combo_hoja.bind("<<ComboboxSelected>>", self._al_cambiar_hoja)

        ttk.Label(superior, text="Tipo de busqueda", style="Panel.TLabel").grid(row=2, column=2, sticky="w", padx=(18, 0), pady=(10, 4))
        self.combo_tipo = ttk.Combobox(
            superior,
            textvariable=self.tipo_busqueda,
            values=["PDF y XML", "Solo PDF"],
            state="readonly",
            width=18,
        )
        self.combo_tipo.grid(row=3, column=2, sticky="w", padx=(18, 0))
        self.combo_tipo.bind("<<ComboboxSelected>>", self._actualizar_controles)

        ttk.Checkbutton(
            superior,
            text="Filtrar por mes",
            variable=self.filtrar_mes,
            command=self._actualizar_controles,
        ).grid(row=3, column=3, sticky="w", padx=(18, 0))
        ttk.Label(superior, text="Columna del mes", style="Panel.TLabel").grid(row=2, column=3, sticky="w", padx=(18, 0), pady=(10, 4))
        self.entrada_columna_mes = ttk.Entry(superior, textvariable=self.columna_mes, width=10)
        self.entrada_columna_mes.grid(row=3, column=3, sticky="e", padx=(130, 0))
        ttk.Label(superior, text="Mes a procesar", style="Panel.TLabel").grid(row=2, column=4, sticky="w", padx=(18, 0), pady=(10, 4))
        self.entrada_mes_filtro = ttk.Entry(superior, textvariable=self.mes_filtro, width=18)
        self.entrada_mes_filtro.grid(row=3, column=4, sticky="w", padx=(18, 0))

        self.notebook = ttk.Notebook(contenedor)
        self.notebook.grid(row=1, column=0, sticky="nsew", pady=(10, 10))

        self._crear_bloque(
            "Factura", self.usar_factura,
            self.busqueda_factura, self.resultado_factura_pdf, self.resultado_factura_xml,
            self.ruta_documentos_factura, self.ruta_destino_factura,
            tiene_xml=True, clave="factura",
        )
        self._crear_bloque(
            "Proveedores", self.usar_proveedores,
            self.busqueda_proveedores_uuid, self.resultado_proveedores_pdf, self.resultado_proveedores_xml,
            self.ruta_documentos_proveedores, self.ruta_destino_proveedores,
            tiene_xml=True, clave="proveedores",
            es_proveedor=True,
            var_nombre_pdf=self.columna_proveedores_nombre_pdf,
        )
        self._crear_bloque(
            "Complemento", self.usar_complemento,
            self.busqueda_complemento, self.resultado_complemento_pdf, self.resultado_complemento_xml,
            self.ruta_documentos_complemento, self.ruta_destino_complemento,
            tiene_xml=True, clave="complemento",
        )
        self._crear_bloque(
            "Pedimento", self.usar_pedimento,
            self.busqueda_pedimento, self.resultado_pedimento_pdf, None,
            self.ruta_documentos_pedimento, self.ruta_destino_pedimento,
            tiene_xml=False, clave="pedimento",
            var_reprocesar=self.reprocesar_pedimento,
        )
        self._crear_b_pedimentos()
        self._crear_bloque(
            "Proforma", self.usar_proforma,
            self.busqueda_proforma, self.resultado_proforma_pdf, None,
            self.ruta_documentos_proforma, self.ruta_destino_proforma,
            tiene_xml=False, clave="proforma",
        )
        self._crear_bloque(
            "Factura Bodega", self.usar_factura_bodega,
            self.busqueda_factura_bodega, self.resultado_factura_bodega_pdf, None,
            self.ruta_documentos_factura_bodega, self.ruta_destino_factura_bodega,
            tiene_xml=False, clave="factura_bodega", folios_multiples=True,
        )
        self._crear_bloque(
            "Pedimento FB", self.usar_pedimento_fb,
            self.busqueda_pedimento_fb, self.resultado_pedimento_fb_pdf, None,
            self.ruta_documentos_pedimento_fb, self.ruta_destino_pedimento_fb,
            tiene_xml=False, clave="pedimento_fb", folios_multiples=True,
        )
        self._crear_bloque(
            "PO", self.usar_po,
            self.busqueda_po, self.resultado_po_pdf, None,
            self.ruta_documentos_po, self.ruta_destino_po,
            tiene_xml=False, clave="po", es_po=True,
            var_nombre_cliente=self.nombre_cliente_po,
            var_reprocesar=self.reprocesar_po,
        )

        pie = ttk.Frame(contenedor, style="Panel.TFrame")
        pie.grid(row=2, column=0, sticky="nsew")
        pie.columnconfigure(0, weight=1)
        pie.rowconfigure(5, weight=1)

        botones = ttk.Frame(pie, style="Panel.TFrame")
        botones.grid(row=0, column=0, columnspan=2, sticky="ew")
        botones.columnconfigure(0, weight=1)

        self.boton_guardar = ttk.Button(botones, text="Guardar Configuracion", command=self._guardar_config)
        self.boton_guardar.grid(row=0, column=2, sticky="e", padx=(0, 8))

        self.boton_restaurar = ttk.Button(botones, text="Restaurar", command=self._restaurar_config)
        self.boton_restaurar.grid(row=0, column=1, sticky="w", padx=(0, 8))

        self.boton_cancelar = ttk.Button(
            botones, text="Cancelar", command=self._cancelar_proceso, state="disabled",
            style="Danger.TButton",
        )
        self.boton_cancelar.grid(row=0, column=3, sticky="e", padx=(0, 8))

        self.boton_iniciar = ttk.Button(
            botones, text="Iniciar", command=self._lanzar_proceso,
            style="Primary.TButton",
        )
        self.boton_iniciar.grid(row=0, column=4, sticky="e")

        self.barra = ttk.Progressbar(pie, orient="horizontal", mode="determinate")
        self.barra.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(10, 6))

        ttk.Label(pie, textvariable=self.contador, style="Panel.TLabel").grid(row=2, column=0, sticky="w", pady=(0, 4))
        ttk.Label(pie, textvariable=self.progreso_texto, style="Panel.TLabel").grid(row=2, column=1, sticky="e", pady=(0, 4))
        ttk.Label(pie, textvariable=self.estado, wraplength=1080, style="Panel.TLabel").grid(
            row=3, column=0, columnspan=2, sticky="w"
        )

        ttk.Label(pie, text="Terminal del proceso", style="Panel.TLabel").grid(
            row=4, column=0, columnspan=2, sticky="w", pady=(10, 4)
        )

        marco_terminal = ttk.Frame(pie, style="Panel.TFrame")
        marco_terminal.grid(row=5, column=0, columnspan=2, sticky="nsew")
        marco_terminal.columnconfigure(0, weight=1)
        marco_terminal.rowconfigure(0, weight=1)

        self.log_terminal = tk.Text(
            marco_terminal,
            height=10,
            wrap="word",
            state="disabled",
            bg="#111827",
            fg="#FFFFFF",
            insertbackground="#FFFFFF",
            selectbackground="#2563EB",
            selectforeground="#FFFFFF",
            relief="solid",
            borderwidth=1,
            font=("Consolas", 10),
            padx=10,
            pady=8,
        )
        self.log_terminal.grid(row=0, column=0, sticky="nsew")

        scroll_terminal = ttk.Scrollbar(
            marco_terminal, orient="vertical", command=self.log_terminal.yview
        )
        scroll_terminal.grid(row=0, column=1, sticky="ns")
        self.log_terminal.configure(yscrollcommand=scroll_terminal.set)

        ttk.Label(
            pie,
            text=f"Disenado por {CREADO_POR}",
            style="Panel.TLabel",
        ).grid(row=6, column=0, columnspan=2, sticky="e", pady=(6, 0))

        self._agregar_log("Aplicacion lista para iniciar.")

        self._actualizar_hojas_excel(mostrar_error=False)
        self._actualizar_controles()

    def _crear_b_pedimentos(self) -> None:
        tab = ttk.Frame(self.notebook, padding=14, style="Panel.TFrame")
        tab.columnconfigure(0, weight=1)
        self.notebook.add(tab, text="B_Pedimentos")

        ttk.Checkbutton(
            tab,
            text="Activar B_Pedimentos",
            variable=self.usar_b_pedimentos,
            command=self._actualizar_controles,
        ).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 10))

        campos = ttk.Frame(tab, style="Panel.TFrame")
        campos.grid(row=1, column=0, sticky="w", pady=(0, 12))

        ttk.Label(campos, text="Columna UUID", style="Panel.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(campos, text="Columna folio/factura", style="Panel.TLabel").grid(row=0, column=1, sticky="w", padx=(22, 0))
        ttk.Label(campos, text="Columna cliente", style="Panel.TLabel").grid(row=0, column=2, sticky="w", padx=(22, 0))
        self.entrada_bp_uuid = ttk.Combobox(campos, textvariable=self.bp_uuid, width=24, state="readonly")
        self.entrada_bp_uuid.grid(row=1, column=0, sticky="w", pady=(4, 10))
        self.entrada_bp_folio = ttk.Combobox(campos, textvariable=self.bp_folio, width=24, state="readonly")
        self.entrada_bp_folio.grid(row=1, column=1, sticky="w", padx=(22, 0), pady=(4, 10))
        self.entrada_bp_cliente = ttk.Combobox(campos, textvariable=self.bp_cliente, width=24, state="readonly")
        self.entrada_bp_cliente.grid(row=1, column=2, sticky="w", padx=(22, 0), pady=(4, 10))

        ttk.Label(campos, text="Escribir pedimento en", style="Panel.TLabel").grid(row=2, column=0, sticky="w")
        ttk.Label(campos, text="Escribir proforma en", style="Panel.TLabel").grid(row=2, column=1, sticky="w", padx=(22, 0))
        self.entrada_bp_resultado_pedimento = ttk.Combobox(
            campos, textvariable=self.bp_resultado_pedimento, width=24, state="readonly"
        )
        self.entrada_bp_resultado_pedimento.grid(row=3, column=0, sticky="w", pady=(4, 10))
        self.entrada_bp_resultado_proforma = ttk.Combobox(
            campos, textvariable=self.bp_resultado_proforma, width=24, state="readonly"
        )
        self.entrada_bp_resultado_proforma.grid(row=3, column=1, sticky="w", padx=(22, 0), pady=(4, 10))

        ttk.Label(campos, text="Filtro por año", style="Panel.TLabel").grid(
            row=0, column=3, sticky="w", padx=(22, 0)
        )
        self.entrada_bp_filtrar = ttk.Entry(campos, textvariable=self.bp_filtrar_carpetas, width=24)
        self.entrada_bp_filtrar.grid(row=1, column=3, sticky="w", padx=(22, 0), pady=(4, 2))
        ttk.Label(campos, text="Ej. 2025 o 2024,2025", style="Panel.TLabel").grid(
            row=2, column=3, sticky="w", padx=(22, 0)
        )
        self.boton_bp_columnas = ttk.Button(
            campos,
            text="Actualizar columnas",
            command=self._actualizar_columnas_b_pedimentos,
        )
        self.boton_bp_columnas.grid(row=3, column=3, sticky="w", padx=(22, 0), pady=(4, 10))

        ttk.Label(tab, text="Carpeta pedimentos", style="Panel.TLabel").grid(row=2, column=0, sticky="w")
        self.entrada_bp_pedimentos = ttk.Entry(tab, textvariable=self.bp_ruta_pedimentos)
        self.entrada_bp_pedimentos.grid(row=3, column=0, sticky="ew", padx=(0, 8), pady=(4, 12))
        self.boton_bp_pedimentos = ttk.Button(
            tab,
            text="Examinar",
            command=lambda: self._seleccionar_carpeta(self.bp_ruta_pedimentos),
        )
        self.boton_bp_pedimentos.grid(row=3, column=1, sticky="e")

        ttk.Label(tab, text="Documentos/facturas", style="Panel.TLabel").grid(row=4, column=0, sticky="w")
        self.entrada_bp_documentos = ttk.Entry(tab, textvariable=self.bp_ruta_documentos)
        self.entrada_bp_documentos.grid(row=5, column=0, sticky="ew", padx=(0, 8), pady=(4, 12))
        self.boton_bp_documentos = ttk.Button(
            tab,
            text="Examinar",
            command=lambda: self._seleccionar_carpeta(self.bp_ruta_documentos),
        )
        self.boton_bp_documentos.grid(row=5, column=1, sticky="e")

        self._actualizar_columnas_b_pedimentos(mostrar_error=False)

    def _actualizar_columnas_b_pedimentos(self, mostrar_error: bool = True) -> None:
        if not hasattr(self, "entrada_bp_uuid"):
            return
        ruta_excel = self.ruta_excel.get().strip()
        nombre_hoja = self.nombre_hoja.get().strip()
        if not ruta_excel or not nombre_hoja or not Path(ruta_excel).is_file():
            valores = ["UUID", "Folio", "Pedimento encontrado", "Proforma encontrada"]
        else:
            try:
                encabezados = self._leer_encabezados_excel_cache(ruta_excel, nombre_hoja)
                valores = encabezados or ["UUID", "Folio", "Pedimento encontrado", "Proforma encontrada"]
            except Exception as e:
                valores = ["UUID", "Folio", "Pedimento encontrado", "Proforma encontrada"]
                if mostrar_error:
                    self._agregar_log(f"B_Pedimentos: no pude cargar encabezados: {e}")

        for sugerido in ["UUID", "Folio", "Pedimento encontrado", "Proforma encontrada"]:
            if sugerido not in valores:
                valores.append(sugerido)

        for combo in (
            self.entrada_bp_uuid,
            self.entrada_bp_folio,
            self.entrada_bp_cliente,
            self.entrada_bp_resultado_pedimento,
            self.entrada_bp_resultado_proforma,
        ):
            combo["values"] = valores

        def asegurar_valor(variable: tk.StringVar, preferidos: list[str]) -> None:
            actual = variable.get().strip()
            if actual in valores:
                return
            claves = {normalizar_para_busqueda(v): v for v in valores}
            for preferido in preferidos:
                valor = claves.get(normalizar_para_busqueda(preferido))
                if valor:
                    variable.set(valor)
                    return
            if valores:
                variable.set(valores[0])

        asegurar_valor(self.bp_uuid, ["UUID", "FOLIO FISCAL", "UUID CFDI"])
        asegurar_valor(self.bp_folio, ["Folio", "FOLIO FACTURA", "FACTURA", "NO FACTURA", "NUMERO FACTURA"])
        asegurar_valor(self.bp_cliente, ["Cliente", "Nombre cliente", "Customer"])
        asegurar_valor(self.bp_resultado_pedimento, ["Pedimento encontrado"])
        asegurar_valor(self.bp_resultado_proforma, ["Proforma encontrada"])

    def _actualizar_hojas_excel(self, mostrar_error: bool = True) -> None:
        if not hasattr(self, "combo_hoja"):
            return

        ruta_excel = self.ruta_excel.get().strip()
        hojas: list[str] = []
        if ruta_excel and Path(ruta_excel).is_file():
            try:
                hojas = self._leer_hojas_excel_cache(ruta_excel)
            except Exception as e:
                if mostrar_error:
                    self._agregar_log(f"No pude leer las hojas del Excel: {e}")

        self.combo_hoja["values"] = hojas
        hoja_actual = self.nombre_hoja.get().strip()
        if hojas:
            if hoja_actual not in hojas:
                self.nombre_hoja.set(hojas[0])
                hoja_actual = hojas[0]
            self.combo_hoja.configure(state="readonly")
            if mostrar_error:
                self._agregar_log(
                    f"Archivo Excel cargado. Hoja seleccionada para trabajar: '{hoja_actual}'."
                )
        else:
            self.nombre_hoja.set("")
            self.combo_hoja.configure(state="disabled")
            if mostrar_error and ruta_excel:
                self._agregar_log("No encontre hojas disponibles en el archivo seleccionado.")

        self._actualizar_columnas_b_pedimentos(mostrar_error=False)

    def _refrescar_hojas_excel(self) -> None:
        ruta_excel = self.ruta_excel.get().strip()
        if ruta_excel and Path(ruta_excel).is_file():
            self._actualizar_hojas_excel(mostrar_error=False)

    def _al_cambiar_ruta_excel(self, _evento=None) -> None:
        self._actualizar_hojas_excel(mostrar_error=True)

    def _al_cambiar_hoja(self, _evento=None) -> None:
        self._actualizar_columnas_b_pedimentos(mostrar_error=False)
        hoja = self.nombre_hoja.get().strip()
        if hoja:
            self.estado.set(f"Hoja seleccionada para trabajar: {hoja}.")
            self._agregar_log(f"Hoja seleccionada: '{hoja}'.")

    def _crear_bloque(
        self,
        nombre: str,
        variable_activa: tk.BooleanVar,
        var_busqueda: tk.StringVar,
        var_resultado_pdf: tk.StringVar,
        var_resultado_xml: Optional[tk.StringVar],
        var_documentos: tk.StringVar,
        var_destino: tk.StringVar,
        tiene_xml: bool,
        clave: str,
        es_po: bool = False,
        folios_multiples: bool = False,
        var_nombre_cliente: Optional[tk.StringVar] = None,
        var_uuid: Optional[tk.StringVar] = None,
        var_resultado_uuid: Optional[tk.StringVar] = None,
        var_reprocesar: Optional[tk.BooleanVar] = None,
        es_proveedor: bool = False,
        var_nombre_pdf: Optional[tk.StringVar] = None,
    ) -> None:
        tab = ttk.Frame(self.notebook, padding=14, style="Panel.TFrame")
        tab.columnconfigure(0, weight=1)
        tab.columnconfigure(1, weight=0)
        self.notebook.add(tab, text=nombre)

        ttk.Checkbutton(
            tab, text=f"Activar {nombre}", variable=variable_activa,
            command=self._actualizar_controles,
        ).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 10))

        campos = ttk.Frame(tab, style="Panel.TFrame")
        campos.grid(row=1, column=0, columnspan=2, sticky="w", pady=(0, 12))

        ttk.Label(campos, text="Buscar en", style="Panel.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(campos, text="Resultado PDF", style="Panel.TLabel").grid(row=0, column=1, sticky="w", padx=(18, 0))

        etiqueta_xml = None
        etiqueta_cliente = None
        etiqueta_uuid = None
        etiqueta_resultado_uuid = None
        entrada_cliente = None
        entrada_uuid = None
        entrada_resultado_uuid = None
        entrada_nombre_pdf = None
        etiqueta_nombre_pdf = None
        if tiene_xml:
            etiqueta_xml = ttk.Label(campos, text="Resultado XML", style="Panel.TLabel")
            etiqueta_xml.grid(row=0, column=2, sticky="w", padx=(18, 0))
            if es_proveedor and var_nombre_pdf is not None:
                etiqueta_nombre_pdf = ttk.Label(campos, text="Renombrar PDF con columna", style="Panel.TLabel")
                etiqueta_nombre_pdf.grid(row=0, column=3, sticky="w", padx=(18, 0))
        elif es_po and var_nombre_cliente is not None:
            etiqueta_cliente = ttk.Label(campos, text="Nombre cliente en", style="Panel.TLabel")
            etiqueta_cliente.grid(row=0, column=2, sticky="w", padx=(18, 0))
        elif var_uuid is not None:
            etiqueta_uuid = ttk.Label(campos, text="UUID en", style="Panel.TLabel")
            etiqueta_uuid.grid(row=0, column=2, sticky="w", padx=(18, 0))
            etiqueta_resultado_uuid = ttk.Label(campos, text="Pedimento UUID en", style="Panel.TLabel")
            etiqueta_resultado_uuid.grid(row=0, column=3, sticky="w", padx=(18, 0))

        entrada_busqueda = ttk.Entry(campos, textvariable=var_busqueda, width=12)
        entrada_busqueda.grid(row=1, column=0, sticky="w", pady=(4, 0))

        entrada_pdf = ttk.Entry(campos, textvariable=var_resultado_pdf, width=12)
        entrada_pdf.grid(row=1, column=1, sticky="w", padx=(18, 0), pady=(4, 0))

        entrada_xml = None
        if tiene_xml and var_resultado_xml is not None:
            entrada_xml = ttk.Entry(campos, textvariable=var_resultado_xml, width=12)
            entrada_xml.grid(row=1, column=2, sticky="w", padx=(18, 0), pady=(4, 0))
            if es_proveedor and var_nombre_pdf is not None:
                entrada_nombre_pdf = ttk.Entry(campos, textvariable=var_nombre_pdf, width=18)
                entrada_nombre_pdf.grid(row=1, column=3, sticky="w", padx=(18, 0), pady=(4, 0))
        elif es_po and var_nombre_cliente is not None:
            entrada_cliente = ttk.Entry(campos, textvariable=var_nombre_cliente, width=12)
            entrada_cliente.grid(row=1, column=2, sticky="w", padx=(18, 0), pady=(4, 0))
        elif var_uuid is not None and var_resultado_uuid is not None:
            entrada_uuid = ttk.Entry(campos, textvariable=var_uuid, width=12)
            entrada_uuid.grid(row=1, column=2, sticky="w", padx=(18, 0), pady=(4, 0))
            entrada_resultado_uuid = ttk.Entry(campos, textvariable=var_resultado_uuid, width=12)
            entrada_resultado_uuid.grid(row=1, column=3, sticky="w", padx=(18, 0), pady=(4, 0))

        inicio_rutas = 2
        check_ocr = None
        check_reprocesar = None
        if es_po:
            ttk.Label(
                tab,
                text="La busqueda de PO revisa nombre, texto interno y OCR solo en PDFs sin texto suficiente.",
                anchor="w", justify="left",
                style="Panel.TLabel",
            ).grid(row=2, column=0, columnspan=2, sticky="w", pady=(0, 10))
            check_ocr = ttk.Checkbutton(
                tab,
                text="Activar OCR texto/imagen para PO",
                variable=self.ocr_po_habilitado,
                command=self._actualizar_controles,
            )
            check_ocr.grid(row=3, column=0, columnspan=2, sticky="w", pady=(0, 10))
            inicio_rutas = 4
        elif folios_multiples:
            ttk.Label(
                tab,
                text="Esta busqueda acepta varios folios en una celda separados por coma.",
                anchor="w", justify="left",
                style="Panel.TLabel",
            ).grid(row=2, column=0, columnspan=2, sticky="w", pady=(0, 10))
            inicio_rutas = 3

        if var_reprocesar is not None:
            check_reprocesar = ttk.Checkbutton(
                tab,
                text="Reprocesar PDF NO",
                variable=var_reprocesar,
                command=self._actualizar_controles,
            )
            check_reprocesar.grid(row=inicio_rutas, column=0, columnspan=2, sticky="w", pady=(0, 10))
            inicio_rutas += 1

        ttk.Label(tab, text="Carpeta de documentos", style="Panel.TLabel").grid(
            row=inicio_rutas, column=0, sticky="w"
        )
        entrada_docs = ttk.Entry(tab, textvariable=var_documentos)
        entrada_docs.grid(
            row=inicio_rutas + 1, column=0, sticky="ew", padx=(0, 8), pady=(4, 12)
        )
        ttk.Button(
            tab, text="Examinar",
            command=lambda v=var_documentos: self._seleccionar_carpeta(v),
        ).grid(row=inicio_rutas + 1, column=1, sticky="e")

        ttk.Label(tab, text="Carpeta destino", style="Panel.TLabel").grid(
            row=inicio_rutas + 2, column=0, sticky="w"
        )
        entrada_destino = ttk.Entry(tab, textvariable=var_destino)
        entrada_destino.grid(
            row=inicio_rutas + 3, column=0, sticky="ew", padx=(0, 8), pady=(4, 0)
        )
        ttk.Button(
            tab, text="Examinar",
            command=lambda v=var_destino: self._seleccionar_carpeta(v),
        ).grid(row=inicio_rutas + 3, column=1, sticky="e")

        self.controles_bloques[clave] = {
            "activa": variable_activa,
            "entrada_busqueda": entrada_busqueda,
            "entrada_pdf": entrada_pdf,
            "entrada_xml": entrada_xml,
            "etiqueta_xml": etiqueta_xml,
            "entrada_cliente": entrada_cliente,
            "etiqueta_cliente": etiqueta_cliente,
            "entrada_uuid": entrada_uuid,
            "entrada_resultado_uuid": entrada_resultado_uuid,
            "etiqueta_uuid": etiqueta_uuid,
            "etiqueta_resultado_uuid": etiqueta_resultado_uuid,
            "entrada_nombre_pdf": entrada_nombre_pdf,
            "etiqueta_nombre_pdf": etiqueta_nombre_pdf,
            "entrada_docs": entrada_docs,
            "entrada_destino": entrada_destino,
            "tiene_xml": tiene_xml,
            "check_ocr": check_ocr,
            "check_reprocesar": check_reprocesar,
        }

    # ------------------------------------------------------------------
    # Control de estado de widgets
    # ------------------------------------------------------------------

    @staticmethod
    def _set_state(widget, habilitado: bool) -> None:
        if widget is not None:
            widget.state(["!disabled"] if habilitado else ["disabled"])

    def _actualizar_controles(self, _evento=None) -> None:
        solo_pdf = self.tipo_busqueda.get() == "Solo PDF"
        filtrar_mes = self.filtrar_mes.get()
        self._set_state(self.entrada_columna_mes, filtrar_mes)
        self._set_state(self.entrada_mes_filtro, filtrar_mes)
        usar_bp = self.usar_b_pedimentos.get()
        for widget in (
            self.entrada_bp_uuid,
            self.entrada_bp_folio,
            self.entrada_bp_cliente,
            self.entrada_bp_resultado_pedimento,
            self.entrada_bp_resultado_proforma,
            self.entrada_bp_pedimentos,
            self.entrada_bp_documentos,
            self.entrada_bp_filtrar,
            self.boton_bp_columnas,
            self.boton_bp_pedimentos,
            self.boton_bp_documentos,
        ):
            self._set_state(widget, usar_bp)
        for ctrl in self.controles_bloques.values():
            activa = ctrl["activa"].get()
            self._set_state(ctrl["entrada_busqueda"], activa)
            self._set_state(ctrl["entrada_pdf"], activa)
            self._set_state(ctrl.get("entrada_cliente"), activa)
            self._set_state(ctrl.get("entrada_uuid"), activa)
            self._set_state(ctrl.get("entrada_resultado_uuid"), activa)
            self._set_state(ctrl.get("entrada_nombre_pdf"), activa)
            self._set_state(ctrl["entrada_docs"], activa)
            self._set_state(ctrl["entrada_destino"], activa)
            self._set_state(ctrl.get("check_ocr"), activa)
            self._set_state(ctrl.get("check_reprocesar"), activa)
            if ctrl["tiene_xml"]:
                self._set_state(ctrl["entrada_xml"], activa and not solo_pdf)
                ctrl["etiqueta_xml"].configure(
                    text="Resultado XML" if not solo_pdf else "Resultado XML (omitido)"
                )

    # ------------------------------------------------------------------
    # Config helpers
    # ------------------------------------------------------------------

    def _obtener_config_actual(self) -> dict:
        return {
            "ruta_excel": self.ruta_excel.get().strip(),
            "nombre_hoja": self.nombre_hoja.get().strip(),
            "tipo_busqueda": self.tipo_busqueda.get().strip(),
            "filtrar_mes": self.filtrar_mes.get(),
            "columna_mes": self.columna_mes.get().strip().upper(),
            "mes_filtro": self.mes_filtro.get().strip(),
            "usar_factura": self.usar_factura.get(),
            "usar_proveedores": self.usar_proveedores.get(),
            "usar_complemento": self.usar_complemento.get(),
            "usar_pedimento": self.usar_pedimento.get(),
            "usar_proforma": self.usar_proforma.get(),
            "usar_factura_bodega": self.usar_factura_bodega.get(),
            "usar_pedimento_fb": self.usar_pedimento_fb.get(),
            "usar_po": self.usar_po.get(),
            "usar_b_pedimentos": self.usar_b_pedimentos.get(),
            "reprocesar_pedimento": self.reprocesar_pedimento.get(),
            "reprocesar_po": self.reprocesar_po.get(),
            "ocr_po_habilitado": self.ocr_po_habilitado.get(),
            "busqueda_factura": self.busqueda_factura.get().strip().upper(),
            "resultado_factura_pdf": self.resultado_factura_pdf.get().strip().upper(),
            "resultado_factura_xml": self.resultado_factura_xml.get().strip().upper(),
            "ruta_documentos_factura": self.ruta_documentos_factura.get().strip(),
            "ruta_destino_factura": self.ruta_destino_factura.get().strip(),
            "busqueda_proveedores_uuid": self.busqueda_proveedores_uuid.get().strip(),
            "resultado_proveedores_pdf": self.resultado_proveedores_pdf.get().strip().upper(),
            "resultado_proveedores_xml": self.resultado_proveedores_xml.get().strip().upper(),
            "columna_proveedores_nombre_pdf": self.columna_proveedores_nombre_pdf.get().strip(),
            "ruta_documentos_proveedores": self.ruta_documentos_proveedores.get().strip(),
            "ruta_destino_proveedores": self.ruta_destino_proveedores.get().strip(),
            "busqueda_complemento": self.busqueda_complemento.get().strip().upper(),
            "resultado_complemento_pdf": self.resultado_complemento_pdf.get().strip().upper(),
            "resultado_complemento_xml": self.resultado_complemento_xml.get().strip().upper(),
            "ruta_documentos_complemento": self.ruta_documentos_complemento.get().strip(),
            "ruta_destino_complemento": self.ruta_destino_complemento.get().strip(),
            "busqueda_pedimento": self.busqueda_pedimento.get().strip().upper(),
            "resultado_pedimento_pdf": self.resultado_pedimento_pdf.get().strip().upper(),
            "uuid_pedimento": self.uuid_pedimento.get().strip().upper(),
            "resultado_pedimento_uuid": self.resultado_pedimento_uuid.get().strip().upper(),
            "ruta_documentos_pedimento": self.ruta_documentos_pedimento.get().strip(),
            "ruta_destino_pedimento": self.ruta_destino_pedimento.get().strip(),
            "bp_uuid": self.bp_uuid.get().strip(),
            "bp_folio": self.bp_folio.get().strip(),
            "bp_cliente": self.bp_cliente.get().strip(),
            "bp_resultado_pedimento": self.bp_resultado_pedimento.get().strip(),
            "bp_resultado_proforma": self.bp_resultado_proforma.get().strip(),
            "bp_ruta_pedimentos": self.bp_ruta_pedimentos.get().strip(),
            "bp_ruta_documentos": self.bp_ruta_documentos.get().strip(),
            "bp_filtrar_carpetas": self.bp_filtrar_carpetas.get().strip(),
            "busqueda_proforma": self.busqueda_proforma.get().strip().upper(),
            "resultado_proforma_pdf": self.resultado_proforma_pdf.get().strip().upper(),
            "ruta_documentos_proforma": self.ruta_documentos_proforma.get().strip(),
            "ruta_destino_proforma": self.ruta_destino_proforma.get().strip(),
            "busqueda_factura_bodega": self.busqueda_factura_bodega.get().strip().upper(),
            "resultado_factura_bodega_pdf": self.resultado_factura_bodega_pdf.get().strip().upper(),
            "ruta_documentos_factura_bodega": self.ruta_documentos_factura_bodega.get().strip(),
            "ruta_destino_factura_bodega": self.ruta_destino_factura_bodega.get().strip(),
            "busqueda_pedimento_fb": self.busqueda_pedimento_fb.get().strip().upper(),
            "resultado_pedimento_fb_pdf": self.resultado_pedimento_fb_pdf.get().strip().upper(),
            "ruta_documentos_pedimento_fb": self.ruta_documentos_pedimento_fb.get().strip(),
            "ruta_destino_pedimento_fb": self.ruta_destino_pedimento_fb.get().strip(),
            "busqueda_po": self.busqueda_po.get().strip().upper(),
            "resultado_po_pdf": self.resultado_po_pdf.get().strip().upper(),
            "nombre_cliente_po": self.nombre_cliente_po.get().strip().upper(),
            "ruta_documentos_po": self.ruta_documentos_po.get().strip(),
            "ruta_destino_po": self.ruta_destino_po.get().strip(),
        }

    def _guardar_config(self) -> None:
        try:
            guardar_configuracion(self._obtener_config_actual())
            self.estado.set(f"Configuracion guardada en {CONFIG_PATH.name}.")
            self._agregar_log(f"Configuracion guardada en {CONFIG_PATH.name}.")
        except OSError as e:
            self.estado.set("No se pudo guardar la configuracion.")
            self._agregar_log(f"Error al guardar configuracion: {e}")

    def _restaurar_config(self) -> None:
        config = dict(CONFIG_DEFAULT)
        for key, val in config.items():
            var = getattr(self, key, None)
            if var is not None:
                var.set(val)
        self._actualizar_hojas_excel(mostrar_error=False)
        self._actualizar_controles()
        guardar_configuracion(config)
        self.estado.set("Configuracion restaurada a valores predeterminados.")
        self._agregar_log("Configuracion restaurada a valores predeterminados.")

    # ------------------------------------------------------------------
    # Dialogos de archivo
    # ------------------------------------------------------------------

    def _seleccionar_excel(self) -> None:
        ruta = filedialog.askopenfilename(
            title="Selecciona el archivo Excel",
            filetypes=[("Archivos Excel", "*.xlsx *.xlsm *.xls")],
        )
        if ruta:
            self.ruta_excel.set(ruta)
            self._actualizar_hojas_excel(mostrar_error=True)

    def _seleccionar_carpeta(self, variable: tk.StringVar) -> None:
        ruta = filedialog.askdirectory(title="Selecciona la carpeta")
        if ruta:
            variable.set(ruta)

    # ------------------------------------------------------------------
    # Construccion de bloques activos
    # ------------------------------------------------------------------

    def _construir_bloques(self, config: dict, tipo_busqueda: str) -> list[BloqueConfig]:
        _ = tipo_busqueda
        definiciones = [
            dict(
                nombre="Factura",
                activo=config["usar_factura"],
                col_busqueda=config["busqueda_factura"],
                col_resultado_pdf=config["resultado_factura_pdf"],
                col_resultado_xml=config["resultado_factura_xml"],
                ruta_documentos=config["ruta_documentos_factura"],
                ruta_destino=config["ruta_destino_factura"],
                tiene_xml=True,
            ),
            dict(
                nombre="Proveedores",
                activo=config.get("usar_proveedores", False),
                col_busqueda=config.get("busqueda_proveedores_uuid", ""),
                col_resultado_pdf=config.get("resultado_proveedores_pdf", ""),
                col_resultado_xml=config.get("resultado_proveedores_xml", ""),
                ruta_documentos=config.get("ruta_documentos_proveedores", ""),
                ruta_destino=config.get("ruta_destino_proveedores", ""),
                tiene_xml=True,
                es_proveedor=True,
                col_nombre_pdf=config.get(
                    "columna_proveedores_nombre_pdf",
                    config.get("columna_proveedores_carpeta", ""),
                ),
            ),
            dict(
                nombre="Complemento",
                activo=config["usar_complemento"],
                col_busqueda=config["busqueda_complemento"],
                col_resultado_pdf=config["resultado_complemento_pdf"],
                col_resultado_xml=config["resultado_complemento_xml"],
                ruta_documentos=config["ruta_documentos_complemento"],
                ruta_destino=config["ruta_destino_complemento"],
                tiene_xml=True,
            ),
            dict(
                nombre="Pedimento",
                activo=config["usar_pedimento"],
                col_busqueda=config["busqueda_pedimento"],
                col_resultado_pdf=config["resultado_pedimento_pdf"],
                ruta_documentos=config["ruta_documentos_pedimento"],
                ruta_destino=config["ruta_destino_pedimento"],
            ),
            dict(
                nombre="Proforma",
                activo=config.get("usar_proforma", False),
                col_busqueda=config.get("busqueda_proforma", ""),
                col_resultado_pdf=config.get("resultado_proforma_pdf", ""),
                ruta_documentos=config.get("ruta_documentos_proforma", ""),
                ruta_destino=config.get("ruta_destino_proforma", ""),
            ),
            dict(
                nombre="Factura Bodega",
                activo=config.get("usar_factura_bodega", False),
                col_busqueda=config.get("busqueda_factura_bodega", ""),
                col_resultado_pdf=config.get("resultado_factura_bodega_pdf", ""),
                ruta_documentos=config.get("ruta_documentos_factura_bodega", ""),
                ruta_destino=config.get("ruta_destino_factura_bodega", ""),
                folios_multiples=True,
            ),
            dict(
                nombre="Pedimento FB",
                activo=config.get("usar_pedimento_fb", False),
                col_busqueda=config.get("busqueda_pedimento_fb", ""),
                col_resultado_pdf=config.get("resultado_pedimento_fb_pdf", ""),
                ruta_documentos=config.get("ruta_documentos_pedimento_fb", ""),
                ruta_destino=config.get("ruta_destino_pedimento_fb", ""),
                folios_multiples=True,
            ),
            dict(
                nombre="PO",
                activo=config["usar_po"],
                col_busqueda=config["busqueda_po"],
                col_resultado_pdf=config["resultado_po_pdf"],
                ruta_documentos=config["ruta_documentos_po"],
                ruta_destino=config["ruta_destino_po"],
                es_po=True,
            ),
        ]
        return [BloqueConfig(**d) for d in definiciones]

    # ------------------------------------------------------------------
    # Lanzar proceso en hilo separado
    # ------------------------------------------------------------------

    def _lanzar_proceso(self) -> None:
        if self.procesando:
            return
        try:
            config_actual = self._obtener_config_actual()
            ruta_excel = config_actual["ruta_excel"]
            nombre_hoja = config_actual["nombre_hoja"]
            tipo_busqueda = config_actual["tipo_busqueda"]

            if not ruta_excel or not Path(ruta_excel).is_file():
                raise ValueError("Debes seleccionar un archivo Excel valido.")
            if not nombre_hoja:
                raise ValueError("Debes seleccionar la hoja de Excel que se va a trabajar.")
            self._validar_hoja_excel(ruta_excel, nombre_hoja)
            if tipo_busqueda not in {"PDF y XML", "Solo PDF"}:
                raise ValueError("Tipo de busqueda no valido.")
            if config_actual.get("filtrar_mes", False):
                if not config_actual.get("columna_mes"):
                    raise ValueError("Debes indicar la columna donde viene el mes.")
                if not config_actual.get("mes_filtro"):
                    raise ValueError("Debes indicar el mes a procesar.")
                letra_a_indice_columna(config_actual["columna_mes"])
                if not claves_mes(config_actual["mes_filtro"]):
                    raise ValueError("El mes a procesar no es valido. Usa por ejemplo Enero, 01 o 2026-01.")

            bloques = self._construir_bloques(config_actual, tipo_busqueda)
            if any(b.activo for b in bloques):
                validar_bloques(bloques, tipo_busqueda, self.cache_indices_origen)
            elif not config_actual.get("usar_b_pedimentos", False):
                raise ValueError("Debes activar al menos un tipo de busqueda.")

            if config_actual.get("usar_b_pedimentos", False):
                if not config_actual.get("bp_uuid") and not config_actual.get("bp_folio"):
                    raise ValueError("En B_Pedimentos debes indicar columna UUID o columna folio/factura.")
                if not config_actual.get("bp_resultado_pedimento") or not config_actual.get("bp_resultado_proforma"):
                    raise ValueError("En B_Pedimentos debes indicar donde escribir pedimento y proforma.")
                if not config_actual.get("bp_ruta_pedimentos") or not Path(config_actual["bp_ruta_pedimentos"]).is_dir():
                    raise ValueError("En B_Pedimentos debes seleccionar la carpeta de pedimentos.")
            guardar_configuracion(config_actual)

        except ValueError as e:
            self.estado.set(f"Error de validacion: {e}")
            self._agregar_log(f"Error de validacion: {e}")
            return

        self.procesando = True
        self._cancelar.clear()
        self._bloquear_botones(True)
        self.estado.set("Iniciando proceso...")
        self._agregar_log("Iniciando proceso...")
        self.contador.set("0 de 0 filas procesadas")
        self.progreso_texto.set("0% | Transcurrido: 00:00 | Restante: --:--")
        self.barra["value"] = 0

        self._hilo = threading.Thread(
            target=self._ejecutar_proceso,
            args=(config_actual, bloques),
            daemon=True,
        )
        self._hilo.start()

    def _cancelar_proceso(self) -> None:
        self._cancelar.set()
        self.estado.set("Cancelando... espera a que termine la fila actual.")
        self._agregar_log("Cancelacion solicitada por el usuario.")

    def _bloquear_botones(self, bloqueado: bool) -> None:
        estado_principal = "disabled" if bloqueado else "!disabled"
        estado_cancelar = "!disabled" if bloqueado else "disabled"
        self.boton_iniciar.state([estado_principal])
        self.boton_guardar.state([estado_principal])
        self.boton_restaurar.state([estado_principal])
        self.boton_cancelar.state([estado_cancelar])

    def _procesar_b_pedimentos(
        self,
        config_actual: dict,
        df: pd.DataFrame,
        hoja_excel,
        filas_permitidas: set[int],
    ) -> None:
        if not config_actual.get("usar_b_pedimentos", False):
            return

        idx_uuid = indice_columna_por_letra_o_encabezado(df, config_actual.get("bp_uuid", ""))
        idx_folio = indice_columna_por_letra_o_encabezado(df, config_actual.get("bp_folio", ""))
        idx_cliente = indice_columna_por_letra_o_encabezado(df, config_actual.get("bp_cliente", ""))
        if idx_uuid is None and idx_folio is None:
            raise ValueError("B_Pedimentos no encontro columna UUID ni columna folio/factura.")

        col_pedimento = indice_hoja_por_letra_o_encabezado(
            hoja_excel,
            config_actual["bp_resultado_pedimento"],
            crear=True,
        )
        col_proforma = indice_hoja_por_letra_o_encabezado(
            hoja_excel,
            config_actual["bp_resultado_proforma"],
            crear=True,
        )

        self._ui(self.estado.set, "B_Pedimentos: indexando PDFs...")
        self._ui(self._agregar_log, "B_Pedimentos: indexando carpeta de pedimentos...")
        pdfs_pedimentos, _ = self._construir_indices_origen(config_actual["bp_ruta_pedimentos"])
        pdfs_pedimentos = filtrar_archivos_por_ruta(
            pdfs_pedimentos,
            config_actual.get("bp_filtrar_carpetas", ""),
        )
        if not pdfs_pedimentos:
            raise ValueError("B_Pedimentos no encontro PDFs en la carpeta de pedimentos.")

        pdfs_documentos = []
        ruta_documentos = config_actual.get("bp_ruta_documentos", "")
        if ruta_documentos:
            if not Path(ruta_documentos).is_dir():
                raise ValueError("La carpeta de documentos/facturas de B_Pedimentos no existe.")
            self._ui(self._agregar_log, "B_Pedimentos: indexando documentos/facturas...")
            pdfs_documentos, _ = self._construir_indices_origen(ruta_documentos)
            pdfs_documentos = filtrar_archivos_por_ruta(
                pdfs_documentos,
                config_actual.get("bp_filtrar_carpetas", ""),
            )

        pdfs_proforma = list(pdfs_documentos)
        rutas_vistas = {ruta for _, _, ruta in pdfs_proforma}
        for pdf in pdfs_pedimentos:
            if pdf[2] not in rutas_vistas:
                pdfs_proforma.append(pdf)

        cache_bp = self.cache_pdf_texto
        encontrados = 0
        faltantes = 0
        filas_objetivo = sorted(filas_permitidas)
        total_bp = len(filas_objetivo)
        self._ui(lambda: self.barra.configure(maximum=max(total_bp, 1), value=0))
        self._ui(self.contador.set, f"B_Pedimentos: 0 de {total_bp} filas revisadas")
        self._ui(
            self._agregar_log,
            f"B_Pedimentos: buscando en {total_bp} filas, {len(pdfs_pedimentos)} PDFs de pedimentos.",
        )
        if idx_cliente is not None:
            self._ui(
                self._agregar_log,
                "B_Pedimentos: columna cliente activa; se priorizaran PDFs por ruta/nombre similar al cliente.",
            )

        for avance, fila_excel in enumerate(filas_objetivo, start=1):
            if self._cancelar.is_set():
                break

            fila_df = fila_excel - 2
            uuid = normalizar_uuid(df.iat[fila_df, idx_uuid]) if idx_uuid is not None else ""
            folio = normalizar_texto(df.iat[fila_df, idx_folio]) if idx_folio is not None else ""
            cliente = normalizar_texto(df.iat[fila_df, idx_cliente]) if idx_cliente is not None else ""
            pdfs_pedimentos_fila = (
                ordenar_pdfs_para_folio_cliente(pdfs_pedimentos, cliente)
                if cliente else pdfs_pedimentos
            )
            pdfs_proforma_fila = (
                ordenar_pdfs_para_folio_cliente(pdfs_proforma, cliente)
                if cliente else pdfs_proforma
            )
            pedimento = ""
            proforma = ""
            metodo = ""

            if uuid:
                ruta_pdf, pedimento = buscar_pedimento_por_uuid_en_pdfs(
                    uuid,
                    pdfs_pedimentos_fila,
                    cache_bp,
                )
                if ruta_pdf:
                    metodo = "UUID/cliente" if cliente else "UUID"

            if not pedimento and folio:
                proforma = buscar_proforma_por_folio_en_pdfs(
                    folio,
                    pdfs_proforma_fila,
                    cache_bp,
                )
                if proforma:
                    ruta_pdf, pedimento = buscar_pedimento_por_referencia_en_pdfs(
                        proforma,
                        pdfs_pedimentos_fila,
                        cache_bp,
                    )
                    if ruta_pdf:
                        metodo = "Folio/proforma/cliente" if cliente else "Folio/proforma"

            hoja_excel.cell(row=fila_excel, column=col_pedimento).value = pedimento or "Falta"
            hoja_excel.cell(row=fila_excel, column=col_proforma).value = proforma or "Falta"

            if pedimento:
                encontrados += 1
                self._ui(
                    self._agregar_log,
                    f"B_Pedimentos fila {fila_excel}: pedimento {pedimento} encontrado por {metodo}.",
                )
            else:
                faltantes += 1

            if avance == 1 or avance % 10 == 0 or avance == total_bp:
                self._ui(lambda n=avance: self.barra.configure(value=n))
                self._ui(self.contador.set, f"B_Pedimentos: {avance} de {total_bp} filas revisadas")
                self._ui(
                    self.estado.set,
                    f"B_Pedimentos: {avance} de {total_bp} filas revisadas.",
                )

        self._ui(
            self._agregar_log,
            f"B_Pedimentos terminado. Encontrados: {encontrados}. Faltan: {faltantes}.",
        )

    # ------------------------------------------------------------------
    # Proceso principal
    # ------------------------------------------------------------------

    def _ejecutar_proceso(self, config_actual: dict, bloques: list[BloqueConfig]) -> None:
        try:
            ruta_excel = config_actual["ruta_excel"]
            nombre_hoja = config_actual["nombre_hoja"]
            tipo_busqueda = config_actual["tipo_busqueda"]
            buscar_xml_general = tipo_busqueda == "PDF y XML"

            self._ui(self.estado.set, f"Leyendo hoja seleccionada: {nombre_hoja}...")
            self._ui(
                self._agregar_log,
                f"Leyendo solo la hoja seleccionada: '{nombre_hoja}'.",
            )
            df = self._leer_dataframe_excel_cache(ruta_excel, nombre_hoja)
            total = len(df.index)
            filas_permitidas = set(range(2, total + 2))
            if config_actual.get("filtrar_mes", False):
                col_mes = config_actual.get("columna_mes", "").strip().upper()
                idx_mes = letra_a_indice_columna(col_mes)
                if idx_mes >= len(df.columns):
                    raise ValueError(f"La columna '{col_mes}' de mes no existe en la hoja.")
                mes_objetivo = config_actual.get("mes_filtro", "").strip()
                filas_permitidas = {
                    fila_excel
                    for fila_excel, valor_mes in enumerate(df.iloc[:, idx_mes].tolist(), start=2)
                    if coincide_mes(valor_mes, mes_objetivo)
                }
                self._ui(
                    self._agregar_log,
                    f"Filtro de mes activo: se procesaran {len(filas_permitidas)} de {total} filas para '{mes_objetivo}'.",
                )
            total_proceso = len(filas_permitidas) if config_actual.get("filtrar_mes", False) else total
            self._ui(lambda: self.barra.configure(maximum=max(total_proceso, 1)))
            self._ui(self.contador.set, f"0 de {total_proceso} filas procesadas")

            indices_archivos: dict[str, tuple] = {}
            indices_destino: dict[str, tuple] = {}
            carpetas_destino: dict[str, Path] = {}
            cache_po = self.cache_po
            cache_texto_completo = self.cache_pdf_texto
            cache_proveedores_texto = self.cache_pdf_texto
            cache_proveedores_ocr = self.cache_pdf_ocr
            resultados_po: dict[str, dict[str, dict]] = {}
            resultados_proveedores: dict[str, dict[str, dict]] = {}
            filas_po_omitidas_por_bloque: dict[str, set[int]] = {}

            for b in bloques:
                if not b.activo:
                    continue
                self._ui(self.estado.set, f"Indexando carpeta: {b.nombre}...")
                self._ui(self._agregar_log, f"Indexando carpeta: {b.nombre}...")
                indices_archivos[b.nombre] = self._construir_indices_origen(b.ruta_documentos)
                carpetas_destino[b.nombre] = Path(b.ruta_destino)
                carpetas_destino[b.nombre].mkdir(parents=True, exist_ok=True)
                indices_destino[b.nombre] = construir_indices(b.ruta_destino)

            libro = load_workbook(ruta_excel)
            hoja_excel = libro[nombre_hoja]

            for b in bloques:
                if not b.activo:
                    continue
                hoja_excel[f"{b.col_resultado_pdf}1"] = f"{b.nombre} Estatus PDF"
                if b.tiene_xml and buscar_xml_general and b.col_resultado_xml:
                    hoja_excel[f"{b.col_resultado_xml}1"] = f"{b.nombre} Estatus XML"

            self._procesar_b_pedimentos(config_actual, df, hoja_excel, filas_permitidas)

            for b in bloques:
                if not b.activo or not b.es_po:
                    continue

                idx = letra_a_indice_columna(b.col_busqueda)
                if idx >= len(df.columns):
                    raise ValueError(
                        f"La columna '{b.col_busqueda}' no existe en la hoja para {b.nombre}."
                    )

                folios_po = [normalizar_texto(valor) for valor in df.iloc[:, idx].tolist()]
                folios_po_busqueda = list(folios_po)
                filas_po_omitidas: set[int] = set()
                pdfs_destino_po, _ = indices_destino[b.nombre]

                self._ui(self._agregar_log, "PO: conciliando Excel contra carpeta destino...")
                conciliados_destino = 0
                for fila_excel_po, folio_po in enumerate(folios_po, start=2):
                    if fila_excel_po not in filas_permitidas:
                        filas_po_omitidas.add(fila_excel_po)
                        folios_po_busqueda[fila_excel_po - 2] = ""
                        continue
                    resultado_destino, _, listo_destino, _ = conciliar_folio_destino(
                        folio_po,
                        pdfs_destino_po,
                        [],
                        False,
                    )
                    hoja_excel[f"{b.col_resultado_pdf}{fila_excel_po}"] = resultado_destino
                    if listo_destino:
                        conciliados_destino += 1
                        filas_po_omitidas.add(fila_excel_po)
                        folios_po_busqueda[fila_excel_po - 2] = ""
                self._ui(
                    self._agregar_log,
                    f"PO: conciliacion destino lista. {conciliados_destino} PDF OK en destino; {total_proceso - conciliados_destino} PDF NO pasan a busqueda.",
                )

                if config_actual.get("reprocesar_po", False):
                    for fila_excel_po in range(2, total + 2):
                        if fila_excel_po not in filas_permitidas:
                            filas_po_omitidas.add(fila_excel_po)
                            folios_po_busqueda[fila_excel_po - 2] = ""
                            continue
                        valor_resultado = hoja_excel[f"{b.col_resultado_pdf}{fila_excel_po}"].value
                        if not resultado_pdf_no(valor_resultado):
                            filas_po_omitidas.add(fila_excel_po)
                            folios_po_busqueda[fila_excel_po - 2] = ""
                    if filas_po_omitidas:
                        self._ui(
                            self._agregar_log,
                            f"PO: modo Reprocesar PDF NO activo; se omiten {len(filas_po_omitidas)} filas que no dicen PDF NO.",
                        )
                filas_po_omitidas_por_bloque[b.nombre] = filas_po_omitidas

                clientes_po: list[str] = []
                col_cliente_po = config_actual.get("nombre_cliente_po", "").strip()
                if col_cliente_po:
                    idx_cliente = letra_a_indice_columna(col_cliente_po)
                    if idx_cliente >= len(df.columns):
                        raise ValueError(
                            f"La columna '{col_cliente_po}' de nombre cliente no existe en la hoja para {b.nombre}."
                        )
                    clientes_po = [
                        normalizar_texto(valor) for valor in df.iloc[:, idx_cliente].tolist()
                    ]

                filas_por_clave: dict[str, list[int]] = {}
                clientes_por_clave: dict[str, str] = {}
                filas_po_resueltas: set[int] = set()

                for fila_excel_po, folio_po in enumerate(folios_po, start=2):
                    clave_po = clave_folio_busqueda(folio_po)
                    if clave_po:
                        filas_por_clave.setdefault(clave_po, []).append(fila_excel_po)
                        if clientes_po:
                            cliente_po = clientes_po[fila_excel_po - 2]
                            if cliente_po and not clientes_por_clave.get(clave_po):
                                clientes_por_clave[clave_po] = cliente_po

                def log_po(mensaje, _self=self):
                    _self._ui(_self._agregar_log, mensaje)

                def actualizar_fila_po(clave, folio, resultado, origen="", _b=b):
                    filas = filas_por_clave.get(clave, [])
                    for fila_excel_po in filas:
                        hoja_excel[f"{_b.col_resultado_pdf}{fila_excel_po}"] = resultado
                        filas_po_resueltas.add(fila_excel_po)

                    total_resueltas = len(filas_po_resueltas)
                    detalle = f" por {origen}" if origen else ""
                    self._ui(lambda n=total_resueltas: self.barra.configure(value=n))
                    self._ui(
                        self.contador.set,
                        f"PO: {total_resueltas} de {total_proceso} filas con resultado",
                    )
                    self._ui(
                        self.estado.set,
                        f"PO: folio '{folio}' {resultado}{detalle}.",
                    )

                self._ui(self.estado.set, f"PO: preparando busqueda en memoria...")
                resultados_po[b.nombre] = preparar_busqueda_po(
                    folios_po_busqueda,
                    indices_archivos[b.nombre][0],
                    cache_po,
                    self.ocr_po_habilitado.get(),
                    clientes_por_clave=clientes_por_clave,
                    log=log_po,
                    on_found=lambda clave, folio, ruta, origen: actualizar_fila_po(
                        clave, folio, "PDF OK", origen
                    ),
                    on_not_found=lambda clave, folio: actualizar_fila_po(
                        clave, folio, "PDF NO"
                    ),
                )

                po_encontradas = sorted(
                    (
                        dato["folio"],
                        dato["ruta"],
                        dato["origen"],
                        clientes_por_clave.get(clave, ""),
                    )
                    for clave, dato in resultados_po[b.nombre].items()
                    if dato["ruta"]
                )
                self._ui(
                    self._agregar_log,
                    f"PO: copiando en lote {len(po_encontradas)} PDFs encontrados...",
                )
                for numero_copia, (folio_po, ruta_pdf, origen_po, cliente_po) in enumerate(po_encontradas, start=1):
                    nombre_copia = f"{folio_po} - {cliente_po}" if cliente_po else folio_po
                    destino_pdf = copiar_documento_como(
                        ruta_pdf,
                        carpetas_destino[b.nombre],
                        nombre_copia,
                    )
                    self._ui(
                        self._agregar_log,
                        f"PO copia {numero_copia}/{len(po_encontradas)}: {nombre_copia} -> {destino_pdf.name} ({origen_po})",
                    )
                self._ui(self._agregar_log, "PO: copia en lote terminada.")

            for b in bloques:
                if not b.activo or not b.es_proveedor:
                    continue

                idx_uuid = indice_columna_por_letra_o_encabezado(df, b.col_busqueda)
                if idx_uuid is None or idx_uuid >= len(df.columns):
                    raise ValueError(
                        f"La columna '{b.col_busqueda}' no existe en la hoja para {b.nombre}."
                    )

                uuids_proveedores = [
                    normalizar_texto(valor)
                    for fila_excel_prov, valor in enumerate(df.iloc[:, idx_uuid].tolist(), start=2)
                    if fila_excel_prov in filas_permitidas
                ]

                def log_proveedores(mensaje, _self=self):
                    _self._ui(_self._agregar_log, mensaje)
                    _self._ui(_self.estado.set, mensaje)

                self._ui(
                    self._agregar_log,
                    f"Proveedores: preparando busqueda en lote para {len(uuids_proveedores)} UUIDs.",
                )
                resultados_proveedores[b.nombre] = preparar_busqueda_proveedores(
                    uuids_proveedores,
                    indices_archivos[b.nombre][0],
                    cache_proveedores_texto,
                    cache_proveedores_ocr,
                    log=log_proveedores,
                )

            fila_excel = 2
            inicio = time.perf_counter()
            folios_revisados: dict[str, set] = {b.nombre: set() for b in bloques}
            filas_procesadas = 0
            iterador_filas = df.itertuples(index=False) if any(b.activo for b in bloques) else []

            for numero, fila in enumerate(iterador_filas, start=1):
                if self._cancelar.is_set():
                    self._ui(self.estado.set, "Proceso cancelado por el usuario.")
                    self._ui(self._agregar_log, "Proceso cancelado por el usuario.")
                    break

                if fila_excel not in filas_permitidas:
                    fila_excel += 1
                    continue

                filas_procesadas += 1
                partes_estado = []

                for b in bloques:
                    if not b.activo:
                        continue

                    if b.es_proveedor:
                        idx = indice_columna_por_letra_o_encabezado(df, b.col_busqueda)
                        if idx is None:
                            raise ValueError(
                                f"La columna '{b.col_busqueda}' no existe en la hoja para {b.nombre}."
                            )
                    else:
                        idx = letra_a_indice_columna(b.col_busqueda)
                    if idx >= len(fila):
                        raise ValueError(
                            f"La columna '{b.col_busqueda}' no existe en la hoja para {b.nombre}."
                        )

                    folio = normalizar_texto(fila[idx])
                    nombre_pdf_proveedor = ""
                    if b.es_proveedor:
                        idx_nombre_pdf = indice_columna_por_letra_o_encabezado(df, b.col_nombre_pdf)
                        if idx_nombre_pdf is None or idx_nombre_pdf >= len(fila):
                            raise ValueError(
                                f"La columna '{b.col_nombre_pdf}' para renombrar PDFs no existe en la hoja para {b.nombre}."
                            )
                        nombre_pdf_proveedor = normalizar_texto(fila[idx_nombre_pdf])
                    usar_xml = buscar_xml_general and b.tiene_xml
                    archivos_pdf, archivos_xml = indices_archivos[b.nombre]
                    archivos_pdf_destino, archivos_xml_destino = indices_destino[b.nombre]
                    folio_log = folio or "(vacio)"
                    omitir_por_filtro_pdf_no = False

                    if b.es_po and fila_excel in filas_po_omitidas_por_bloque.get(b.nombre, set()):
                        omitir_por_filtro_pdf_no = True

                    if omitir_por_filtro_pdf_no:
                        resultado_pdf = normalizar_texto(
                            hoja_excel[f"{b.col_resultado_pdf}{fila_excel}"].value
                        ) or "Sin estatus"
                        resultado_xml = ""
                        self._ui(
                            self._agregar_log,
                            f"Fila {numero}/{total} - {b.nombre}: '{folio_log}' no dice PDF NO, se omite.",
                        )
                        partes_estado.append(f"{b.nombre}: {folio_log} | PDF: {resultado_pdf} (omitido)")
                        continue

                    if not b.es_po:
                        if b.folios_multiples:
                            (
                                resultado_pdf_destino,
                                encontrados_destino,
                                faltantes_destino,
                                listo_destino,
                            ) = conciliar_folios_multiples_destino(
                                folio,
                                archivos_pdf_destino,
                                usar_coincidencia_pedimento=(b.nombre == "Pedimento FB"),
                            )
                            hoja_excel[f"{b.col_resultado_pdf}{fila_excel}"] = resultado_pdf_destino
                            if listo_destino:
                                self._ui(
                                    self._agregar_log,
                                    f"Fila {numero}/{total} - {b.nombre}: todos los folios ya estan en destino: {', '.join(encontrados_destino)}.",
                                )
                                partes_estado.append(f"{b.nombre}: {folio_log} | PDF: PDF OK (destino)")
                                continue
                            if encontrados_destino:
                                self._ui(
                                    self._agregar_log,
                                    f"Fila {numero}/{total} - {b.nombre}: en destino ya estan: {', '.join(encontrados_destino)}.",
                                )
                            if faltantes_destino:
                                self._ui(
                                    self._agregar_log,
                                    f"Fila {numero}/{total} - {b.nombre}: faltan en destino: {', '.join(faltantes_destino)}. Inicia busqueda.",
                                )
                        else:
                            (
                                resultado_pdf_destino,
                                resultado_xml_destino,
                                listo_destino,
                                pdf_en_destino,
                            ) = conciliar_folio_destino(
                                folio,
                                archivos_pdf_destino,
                                archivos_xml_destino,
                                usar_xml,
                                usar_coincidencia_pedimento=(b.nombre == "Pedimento"),
                            )
                            hoja_excel[f"{b.col_resultado_pdf}{fila_excel}"] = resultado_pdf_destino
                            if usar_xml and b.col_resultado_xml:
                                hoja_excel[f"{b.col_resultado_xml}{fila_excel}"] = resultado_xml_destino
                            if listo_destino:
                                self._ui(
                                    self._agregar_log,
                                    f"Fila {numero}/{total} - {b.nombre}: '{folio_log}' ya esta completo en destino.",
                                )
                                parte_destino = f"{b.nombre}: {folio_log} | PDF: PDF OK (destino)"
                                if usar_xml:
                                    parte_destino += " | XML: XML OK (destino)"
                                partes_estado.append(parte_destino)
                                continue
                            self._ui(
                                self._agregar_log,
                                f"Fila {numero}/{total} - {b.nombre}: '{folio_log}' no esta completo en destino. Inicia busqueda.",
                            )

                    if not folio:
                        self._ui(
                            self._agregar_log,
                            f"Fila {numero}/{total} - {b.nombre}: folio vacio, se marca como no encontrado.",
                        )
                    elif b.folios_multiples:
                        folios_multiples = separar_folios(folio)
                        self._ui(
                            self._agregar_log,
                            f"Fila {numero}/{total} - {b.nombre}: buscando {len(folios_multiples)} folios separados por coma.",
                        )
                    elif b.es_po:
                        self._ui(
                            self._agregar_log,
                            f"Fila {numero}/{total} - {b.nombre}: consultando resultado en memoria para '{folio_log}'.",
                        )
                    elif b.es_proveedor:
                        self._ui(
                            self._agregar_log,
                            f"Fila {numero}/{total} - {b.nombre}: consultando busqueda en lote para '{folio_log}'.",
                        )
                    else:
                        detalle_busqueda = "nombre del archivo"
                        self._ui(
                            self._agregar_log,
                            f"Fila {numero}/{total} - {b.nombre}: buscando '{folio_log}' en {detalle_busqueda}...",
                        )

                    if b.es_po:
                        resultado_xml = ""
                        clave_po = clave_folio_busqueda(folio)
                        dato_po = resultados_po.get(b.nombre, {}).get(
                            clave_po, {"ruta": None, "origen": ""}
                        )
                        if not folio:
                            resultado_pdf = "PDF NO"
                        elif clave_po in folios_revisados[b.nombre]:
                            resultado_pdf = "PDF OK"
                        elif dato_po["ruta"]:
                            resultado_pdf = "PDF OK"
                            folios_revisados[b.nombre].add(clave_po)
                        else:
                            resultado_pdf = "PDF NO"
                            folios_revisados[b.nombre].add(clave_po)
                    elif b.folios_multiples:
                        resultado_xml = ""
                        resultado_pdf, encontrados, faltantes, repetidos = procesar_folios_multiples(
                            folio,
                            archivos_pdf,
                            carpetas_destino[b.nombre],
                            folios_revisados[b.nombre],
                            buscar_texto_completo=(b.nombre == "Pedimento FB"),
                            cache_texto=cache_texto_completo,
                        )
                        if encontrados:
                            self._ui(
                                self._agregar_log,
                                f"Fila {numero}/{total} - {b.nombre}: folios encontrados: {', '.join(encontrados)}.",
                            )
                        if faltantes:
                            self._ui(
                                self._agregar_log,
                                f"Fila {numero}/{total} - {b.nombre}: folios no encontrados: {', '.join(faltantes)}.",
                            )
                        if repetidos:
                            self._ui(
                                self._agregar_log,
                                f"Fila {numero}/{total} - {b.nombre}: folios ya revisados marcados como PDF OK: {', '.join(repetidos)}.",
                            )
                    else:
                        if b.es_proveedor:
                            resultado_xml = ""
                            clave_proveedor = normalizar_para_busqueda(folio)
                            dato_proveedor = resultados_proveedores.get(b.nombre, {}).get(
                                clave_proveedor,
                                {"ruta": None, "origen": ""},
                            )
                            ruta_pdf = dato_proveedor.get("ruta")
                            origen_pdf = dato_proveedor.get("origen", "")
                            ruta_xml = existe_documento(folio, archivos_xml) if usar_xml else None
                            if ruta_pdf:
                                nombre_copia = nombre_pdf_proveedor or folio
                                destino_pdf = copiar_documento_como(
                                    ruta_pdf,
                                    carpetas_destino[b.nombre],
                                    nombre_copia,
                                )
                                resultado_pdf = "PDF OK"
                                self._ui(
                                    self._agregar_log,
                                    f"Fila {numero}/{total} - {b.nombre}: UUID '{folio_log}' encontrado por {origen_pdf} y copiado como '{destino_pdf.name}'.",
                                )
                            else:
                                resultado_pdf = "PDF NO"

                            if usar_xml and ruta_xml:
                                copiar_documento(ruta_xml, carpetas_destino[b.nombre])
                                resultado_xml = "XML OK"
                            elif usar_xml:
                                resultado_xml = "XML NO"
                        elif b.nombre == "Pedimento":
                            resultado_pdf = "PDF NO"
                            resultado_xml = ""
                            if not folio:
                                resultado_pdf = "PDF NO"
                            elif clave_folio_busqueda(folio) in folios_revisados[b.nombre]:
                                resultado_pdf = "PDF OK"
                            else:
                                ruta_pdf = existe_documento_pedimento(folio, archivos_pdf)
                                if ruta_pdf:
                                    copiar_documento(ruta_pdf, carpetas_destino[b.nombre])
                                    resultado_pdf = "PDF OK"
                                    folios_revisados[b.nombre].add(clave_folio_busqueda(folio))
                                else:
                                    folios_revisados[b.nombre].add(clave_folio_busqueda(folio))
                        else:
                            resultado_pdf, resultado_xml = procesar_folio(
                                folio,
                                archivos_pdf,
                                archivos_xml,
                                carpetas_destino[b.nombre],
                                folios_revisados[b.nombre],
                                usar_xml,
                            )
                        if (
                            b.nombre == "Pedimento"
                            and resultado_pdf == "PDF NO"
                            and folio
                        ):
                            ruta_pdf_texto = buscar_folio_en_texto_pdfs_completos(
                                folio,
                                archivos_pdf,
                                cache_texto_completo,
                            )
                            if ruta_pdf_texto:
                                copiar_documento_como(
                                    ruta_pdf_texto,
                                    carpetas_destino[b.nombre],
                                    folio,
                                )
                                resultado_pdf = "PDF OK"
                                self._ui(
                                    self._agregar_log,
                                    f"Fila {numero}/{total} - {b.nombre}: '{folio_log}' encontrado en texto de PDF completo y copiado con nombre del folio.",
                                )
                            else:
                                resultado_pdf = f"PDF NO: {folio}"

                    if folio:
                        if normalizar_texto(resultado_pdf).casefold().startswith("pdf ok"):
                            origen_po = ""
                            if b.es_po:
                                origen_po = resultados_po.get(b.nombre, {}).get(
                                    clave_folio_busqueda(folio), {}
                                ).get("origen", "")
                            detalle_origen = f" por {origen_po}" if origen_po else ""
                            self._ui(
                                self._agregar_log,
                                f"Fila {numero}/{total} - {b.nombre}: PDF encontrado{detalle_origen} para '{folio_log}'.",
                            )
                        elif b.es_po:
                            self._ui(
                                self._agregar_log,
                                f"Fila {numero}/{total} - {b.nombre}: PDF no encontrado para '{folio_log}' despues de la busqueda PO completa.",
                            )
                        else:
                            self._ui(
                                self._agregar_log,
                                f"Fila {numero}/{total} - {b.nombre}: PDF no encontrado para '{folio_log}', sigue buscando.",
                            )

                        if usar_xml:
                            if resultado_xml == "XML OK":
                                self._ui(
                                    self._agregar_log,
                                    f"Fila {numero}/{total} - {b.nombre}: XML encontrado para '{folio_log}'.",
                                )
                            elif resultado_xml == "XML NO":
                                self._ui(
                                    self._agregar_log,
                                    f"Fila {numero}/{total} - {b.nombre}: XML no encontrado para '{folio_log}', sigue buscando.",
                                )

                    hoja_excel[f"{b.col_resultado_pdf}{fila_excel}"] = resultado_pdf
                    if usar_xml and b.col_resultado_xml:
                        hoja_excel[f"{b.col_resultado_xml}{fila_excel}"] = resultado_xml

                    etiqueta = b.nombre
                    parte = f"{etiqueta}: {folio_log} | PDF: {resultado_pdf}"
                    if usar_xml:
                        parte += f" | XML: {resultado_xml}"
                    partes_estado.append(parte)

                transcurrido = time.perf_counter() - inicio
                porcentaje = (filas_procesadas / total_proceso * 100) if total_proceso else 0
                restante = (
                    (transcurrido / filas_procesadas) * (total_proceso - filas_procesadas)
                    if filas_procesadas else None
                )

                self._ui(lambda n=filas_procesadas: self.barra.configure(value=n))
                self._ui(self.contador.set, f"{filas_procesadas} de {total_proceso} filas procesadas")
                self._ui(
                    self.progreso_texto.set,
                    f"{porcentaje:.0f}% | Transcurrido: {formatear_segundos(transcurrido)} "
                    f"| Restante: {formatear_segundos(restante)}",
                )
                self._ui(
                    self.estado.set,
                    " | ".join(partes_estado) if partes_estado else "Sin bloques activos.",
                )

                fila_excel += 1

            if not self._cancelar.is_set():
                self._ui(self.estado.set, "Guardando Excel...")
                self._ui(self._agregar_log, "Guardando Excel...")
                ruta_tmp = Path(ruta_excel).with_suffix(".tmp.xlsx")
                libro.save(ruta_tmp)
                ruta_tmp.replace(ruta_excel)
                self._invalidar_cache_excel(ruta_excel)
                self._ui(
                    self.estado.set,
                    "Proceso terminado. El Excel fue actualizado y los archivos copiados.",
                )
                self._ui(
                    self._agregar_log,
                    "Proceso terminado. El Excel fue actualizado y los archivos encontrados fueron copiados.",
                )
            else:
                libro.save(ruta_excel)
                self._invalidar_cache_excel(ruta_excel)
                self._ui(self._agregar_log, "Se guardo el progreso parcial por cancelacion.")

        except Exception as e:
            self._ui(self.estado.set, f"Error: {e}")
            self._ui(self._agregar_log, f"Error: {e}")
        finally:
            self.procesando = False
            self._ui(self._bloquear_botones, False)

    def _ui(self, fn, *args):
        if args:
            self.root.after(0, lambda: fn(*args))
        else:
            self.root.after(0, fn)

    def _agregar_log(self, mensaje: str) -> None:
        if not self.log_terminal:
            return
        marca_tiempo = time.strftime("%H:%M:%S")
        self.log_terminal.configure(state="normal")
        self.log_terminal.insert("end", f"[{marca_tiempo}] {mensaje}\n")
        self.log_terminal.see("end")
        self.log_terminal.configure(state="disabled")

    # ------------------------------------------------------------------
    # Arranque
    # ------------------------------------------------------------------

    def ejecutar(self) -> None:
        self.root.mainloop()


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    app = Aplicacion()
    app.ejecutar()
